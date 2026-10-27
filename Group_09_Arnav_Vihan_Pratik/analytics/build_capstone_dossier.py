"""Condense the faculty form to four pages, retaining all 212 tagged controls."""
from pathlib import Path
from copy import deepcopy
import json,zipfile,hashlib
from lxml import etree as E
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
G=Path(__file__).resolve().parents[1];A=G/'analytics';NS={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
SOURCE=G/'archive/dossier_original_template.docx'

def main():
    root=E.fromstring(zipfile.ZipFile(SOURCE).read('word/document.xml'))
    controls={s.xpath('./w:sdtPr/w:tag/@w:val',namespaces=NS)[0]:s for s in root.xpath('//w:sdt',namespaces=NS)}
    original_props={k:E.tostring(v.find('w:sdtPr',NS),method='c14n') for k,v in controls.items()}
    values={};used=[];stats=json.loads((A/'capstone_500hz/statistics.json').read_text());cmp=stats['comparisons'];v=json.loads((A/'capstone_500hz/verification.json').read_text())
    def setv(prefix,**kwargs):values.update({prefix+'.'+k:str(x) for k,x in kwargs.items()})
    for i in (1,2):setv(f'M09.S1.claim{i}',source='Unmeasured; excluded.',source_type='No source found',verdict='Delete')
    setv('M09.S1.claim3',source='TrialConfig: 0.35 s assumed delay.',source_type='No source found',verdict='Correct')
    for key,app,why in [('iso10218','No','No industrial arm integration.'),('iso3691','Partly','ISO 3691-4:2023 scope concepts only; no conformity claim.'),('ria','No','Industrial arm standard; outside this evaluation.'),('iec60601','No','Not medical electrical equipment.'),('domain','Partly','SI, x-forward/y-left/z-up; rigid yard; no personnel or explosive-hazard validation.')]:setv('M09.S1.std_'+key,applies=app,clause=why)
    setv('M09.S1',baseline_value='Reactive driver: 0/30 accepted; goal error 12.310 ± 0.093 m.',baseline_source='Own simulation: analytics/capstone_500hz/statistics.json; no field baseline.',stakeholder='Inspection/response staff; exposure and workload unmeasured.')
    setv('M09.S2',timestep='0.002 s; actual 500 Hz physics, torque and CSV.',friction='Robot 1/0.01/0.001; floor sliding U(0.75,1.15), torsional 0.005, rolling 0.0001.',solref='Robot 0.005/1; floor 0.02/1. Numerical compliance; uncalibrated materials.',control_law='Other',gains='Leg PD Kp=500/800/800 N m/rad; Kd=25/40/40 N m s/rad; caps 76.4 N m leg, 21.6 N m hub.',fsm_states='WAKEUP → DETOUR_NAV → DOCK_ALIGN → front pair lift/support → BODY_ADVANCE → rear pair lift/support → SETTLE → CLEAR_REAR_EDGE; repeat.',model_loads='Yes')
    for i,x in enumerate(['M20 anatomy + physical course.','Entry point to measured runner.','Payback, NCRI, labor, Erlang C.','Ideal-ray mapping and A*.','Contact-gated climb; 500 Hz logger.'],1):setv(f'M09.S2.script{i}',change=x)
    setv('M09.S2.script4',name='M20/m20_rosbot_mapping.py');setv('M09.S2.script5',name='M20/simulation.py')
    setv('M09.S3.E073',hreclaimed_inputs='Assume 312 h/year; full-time 2080 h/year.',hreclaimed_metric='H=312 h/year; ΔFTE=0.150.',payback_inputs='Assume K=0.30 annual-baseline units; κ=0.22; life L=5 years.',payback_metric='P=K/[L(1−κ)]=0.076923; T=4.615 months.',ncri_inputs='Assume annual baseline=1, robot=0.22.',ncri_metric='NCRI=1−κ=0.780.',queue_inputs='Assume λ=8/week, μ=12/week/server; c=1; 40 working h/week.',queue_metric='ρ=0.667; Wq=1/6 week=6.667 working h.',sensitivity_inputs='κ=0.12…1.20; labor ±30%; utilization 0…>1.',sensitivity_metric='Life break-even κ=0.94; net savings vanish at 1. FTE=0.105…0.195.',queue_model='M/M/1',queue_params='λ=8/week; μ=12/week; c=1. Poisson, exponential, FCFS assumed.',payback_months='4.615385 months (planning scenario).',reproducible='Yes')
    fmea=[('Tip-over','Mission loss',8,5,6,3,3,'Tilt latch; staged transfer.'),('Delayed stop','Hazard contact',8,6,6,3,2,'Local ray stop and mapped detour.'),('Blocked passage','No route',6,5,6,3,2,'Unknown blocked; no-route stop.'),('Step lip error','Foot scrape',7,6,5,3,2,'Measured lip; paired lift/support gates.'),('Distance diagnostic','False clearance',5,4,7,3,2,'Witness check; solver contact audit.')]
    for i,(mode,effect,s,o,d,o2,d2,mit) in enumerate(fmea,1):
        setv(f'M09.S3.fmea{i}',effect=effect,severity=s,occurrence=o,detection=d,rpn_pre=s*o*d,mitigation=mit,rpn_post=f'{s*o2*d2} ({s}×{o2}×{d2})',evidence='Both')
        if i>2:setv(f'M09.S3.fmea{i}',mode=mode)
    setv('M09.S4',baseline_condition='Delayed reactive software driver; same M20; no map/paired climb. Integration comparison, not single-feature ablation.',n_runs='30 independent runs/policy; 10 at each roughness bound 0, 3, 6 mm.',randomised='Seeds baseline 31000–31029, proposed 32000–32029; floor μ U(.75,1.15); barriers ±.08 m; start y ±.02 m/yaw ±.02 rad; rubble height U(bound/2,bound).',log_rate_hz=f'500 Hz; {v["total_actual_samples"]:,} actual samples; every Δt=0.002 s checked.',seeds_recorded='Yes',real_csv_path='analytics/capstone_500hz/*/*_real_trials.csv; aggregate capstone_real_trials.csv; design/schema/source hashes retained.')
    for key,metric,unit in [('rmse','tracking_rmse_m','m'),('effort','control_effort_Nm2_s','N² m² s'),('metric','goal_error_m','m')]:
        c=cmp[metric];fmt='.0f' if key=='effort' else '.4f'
        setv('M09.S4.'+key,n='30/arm',baseline=f'{c["baseline_mean"]:{fmt}} ± {c["baseline_SD"]:{fmt}}',proposed=f'{c["proposed_mean"]:{fmt}} ± {c["proposed_SD"]:{fmt}}',unit=unit,better='Lower is better')
    setv('M09.S4.metric',name='Final XY goal error')
    c=cmp['goal_error_m'];setv('M09.S4',test="Welch's t-test",t_df_p=f't={c["t"]:.3f}; df={c["df"]:.3f}; p={c["p_two_sided"]:.3g} (two-sided).',cohens_d=f'{c["cohen_d_pooled"]:.3f}, pooled SD; proposed minus baseline.',ci95=f'[{c["CI95"][0]:.4f}, {c["CI95"][1]:.4f}] m; mean difference {c["difference_proposed_minus_baseline"]:.4f} m.')
    for key in ['prov_own_runs','prov_no_simulated','prov_plots_reproducible','prov_economics_sourced']:values['M09.S4.'+key]='Yes'
    setv('M09.S5',linkedin_url='PENDING — team will enter genuine post URL.',video_length='PENDING — target 03:45.',live_mujoco='No',business_graph='No',all_members_speak='No')
    cases=[('Earlier delivery used a car.','Checked the original source','Supplied M20.xml: four legs, 16 motors, 17 meshes.','Other'),('Ray cutoff treated as exact range.','Official documentation (MuJoCo, standard)','mj_multiRay coarse culling differs from hit-range filtering; regression test passes.','Wrong code / MuJoCo API'),('Distance zero accepted without witness check.','Ran / tested the code or simulation','Inconsistent witnesses withheld; use actual solver contacts at 2 ms.','Wrong code / MuJoCo API')]
    for i,(said,caught,truth,err) in enumerate(cases,1):setv(f'M09.S5.ai{i}',tool='Other',said=said,caught_by=caught,truth=truth,error_type=err,verdict='Corrected')
    doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
    sec.top_margin=sec.bottom_margin=Inches(.38);sec.left_margin=sec.right_margin=Inches(.42);sec.header_distance=sec.footer_distance=Inches(.16)
    normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(8.5);normal.paragraph_format.space_after=Pt(2);normal.paragraph_format.line_spacing=1
    for name,size in [('Title',14),('Heading 1',11),('Heading 2',9)]:
        st=doc.styles[name];st.font.name='Calibri';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0);st.paragraph_format.space_after=Pt(3);st.paragraph_format.space_before=Pt(4)
    for st in doc.styles:
        for b in list(st.element.iter(qn('w:pBdr'))):b.getparent().remove(b)
    def p(text,style=None):return doc.add_paragraph(text,style)
    def h(text,page=False):
        x=p(text,'Heading 1');x.paragraph_format.page_break_before=page;return x
    def control(k):
        assert k in controls and k not in used,k;used.append(k);s=deepcopy(controls[k])
        if k in values:
            opts=s.xpath('./w:sdtPr/w:dropDownList/w:listItem/@w:displayText',namespaces=NS)
            assert not opts or values[k] in opts,(k,values[k])
            ts=s.xpath('./w:sdtContent//w:t',namespaces=NS);ts[0].text=values[k]
            for t in ts[1:]:t.text=''
        for r in s.xpath('./w:sdtContent//w:r',namespaces=NS):
            rp=r.find('w:rPr',NS)
            if rp is None:rp=E.SubElement(r,qn('w:rPr'))
            for old in list(rp):
                if old.tag in [qn('w:sz'),qn('w:szCs'),qn('w:color')]:rp.remove(old)
            E.SubElement(rp,qn('w:sz'),{qn('w:val'):'17'});E.SubElement(rp,qn('w:color'),{qn('w:val'):'000000'})
        return s
    def table(headers,rows,widths):
        t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
        for col,width in zip(t.columns,widths):col.width=Inches(width)
        for i,row in enumerate([headers,*rows]):
            cells=t.rows[0].cells if i==0 else t.add_row().cells
            for cell,value,width in zip(cells,row,widths):
                cell.width=Inches(width);par=cell.paragraphs[0];par.paragraph_format.space_after=Pt(1);par.paragraph_format.space_before=Pt(1)
                for j,item in enumerate(value if isinstance(value,list) else [value]):
                    if j:par.add_run('\n')
                    if str(item).startswith('M09.'):par._p.append(control(item))
                    else:par.add_run(str(item)).bold=i==0
                prop=cell._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
                for side in ['top','bottom','left','right']:
                    x=OxmlElement('w:'+side);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'3');x.set(qn('w:color'),'D9D9D9');b.append(x)
                prop.append(b);m=OxmlElement('w:tcMar')
                for side in ['top','bottom','left','right']:
                    x=OxmlElement('w:'+side);x.set(qn('w:w'),'35');x.set(qn('w:type'),'dxa');m.append(x)
                prop.append(m)
                if i==0:
                    x=OxmlElement('w:shd');x.set(qn('w:fill'),'E8EDF0');prop.append(x)
            t.rows[i]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        return t
    def fields(prefix,pairs):table(['Field','Recorded answer'],[(label,prefix+'.'+key) for key,label in pairs],[1.55,5.85])
    def fig(path,width=7.2):
        x=doc.add_paragraph();x.paragraph_format.space_after=Pt(1);x.add_run().add_picture(str(path),width=Inches(width))
    p('GROUP 09 M20 TECHNICAL DOSSIER','Title')
    p("SVKM’s NMIMS MPSTME Mumbai | MDRIIA 702CO0E012 | CSBS | AY 2026–27\nCapstone ICA: Robotics 10 + Business 5 + Showcase and Viva 5 = 20 marks")
    table(['Student','Roll','SAP ID','Assigned review and viva responsibility'],[
        ['Arnav Saurabh Surve','E064','70362400067','Robotics architecture and dynamics'],['Vihan Shripad Joshi','E070','70362400086','Navigation and telemetry QA'],['Pratik Mangesh Gaikwad','E073','70362300032','CSBS business analysis']],[1.7,.5,1.1,4.1])
    h('1 Problem and verified claims')
    p('M20 hybrid quadruped for hazardous-terrain reconnaissance. Research question: can local traversability mapping and paired-leg transfer improve unknown-yard navigation? Human workload and communication savings remain untested. Faculty title motivation is retained; numeric human benefits are excluded.')
    table(['Original claim','Source or correction','Source type','Verdict'],[[x,f'M09.S1.claim{i}.source',f'M09.S1.claim{i}.source_type',f'M09.S1.claim{i}.verdict'] for i,x in enumerate(['80% bandwidth reduction','NASA-TLX 78 → 32','350 ms latency'],1)],[1.5,3.2,1.6,1.1])
    table(['Standard','Applies','Requirement or scope'],[[label,'M09.S1.std_'+key+'.applies','M09.S1.std_'+key+'.clause'] for key,label in [('iso10218','ISO 10218'),('iso3691','ISO 3691-4'),('ria','ANSI/RIA R15.06'),('iec60601','IEC 60601-1'),('domain','Domain assumptions')]],[1.3,.6,5.5])
    fields('M09.S1',[('baseline_value','Baseline'),('baseline_source','Evidence'),('stakeholder','Stakeholder')])
    h('2 Physics and implemented controller')
    p('M20/M20.xml: 34.525856 kg; nq=23, nv=22, nu=16; 12 leg + 4 hub motors; 17 supplied meshes. Industrial yard: three alternating barriers and two 8 cm terrace rises. Wheels cruise; articulated legs climb. This is not demonstrated full-time braked walking.')
    fields('M09.S2',[(x,l) for x,l in [('timestep','Physics and logging'),('friction','Friction'),('solref','Contact solref'),('control_law','Law'),('gains','Gains and limits'),('fsm_states','FSM stages'),('model_loads','Model runs')]])
    table(['File','Authorship','Author','Change'],[[name,f'M09.S2.script{i}.authorship',f'M09.S2.script{i}.author',f'M09.S2.script{i}.change'] for i,name in enumerate(['models/hazardous_terrain_ugv.xml','src/rough_terrain_recon_controller.py','analytics/hazardous_recon_economics.py','M09.S2.script4.name','M09.S2.script5.name'],1)],[2.7,1.2,.65,2.85])
    p('Students attest authorship from genuine commits. Substantial AI assistance; supplied model reused.')
    # Actual excerpt, not pseudocode; explain the remaining system in the annex.
    lines=(G/'M20/recon_controller.py').read_text().splitlines();start=next(i for i,l in enumerate(lines) if l.startswith('def track('))
    snippet='M20/recon_controller.py\n'+'\n'.join(f'{i+1}: {lines[i]}' for i in range(start,min(start+9,len(lines))))
    x=p(snippet);x.paragraph_format.space_after=Pt(1)
    for r in x.runs:r.font.name='Consolas';r.font.size=Pt(8)
    h('2.4 Free body diagram and worked mechanics',page=True)
    p('       ↑ N₁ + N₂                  ↑ N₃ + N₄\n       └──────── M20 body ────────┘    → tangential drive force\n                          ↓ mg')
    p('M(q)q̈+Cq̇+g=τ+Jᵀf. At rest mg=34.525856×9.81=338.699 N; equal-load approximation 84.675 N/foot. Actual seed-42 normal sum=338.739 N at 2.302 s (analytics/capstone_checks/pilot_42/telemetry.csv); normals need not be equal.')
    h('3 CSBS techno economic systems analysis')
    p('Owner E073. All parities, labor, lifetime and queue rates below are explicit uncalibrated planning assumptions. Successful simulated navigation does not establish observed labor savings or safety benefit. Script: analytics/hazardous_recon_economics.py; interactive view: analytics/business_dashboard.py.')
    table(['Quantity and equation','Inputs and provenance','Result'],[[label,'M09.S3.E073.'+key+'_inputs','M09.S3.E073.'+key+'_metric'] for key,label in [('hreclaimed','H; ΔFTE=H/2080'),('payback','P=K/[L(1−κ)]; T=12K/(1−κ)'),('ncri','NCRI=1−κ'),('queue','ρ=λ/(cμ); Wq=ρ/(μ−λ) for c=1'),('sensitivity','Scenario variation and break-even')]],[2.15,2.7,2.55])
    fields('M09.S3.E073',[('queue_model','Queue model'),('queue_params','Queue parameters'),('payback_months','Payback months'),('reproducible','Script reproduces')])
    p('Queue assumptions: Poisson arrivals, independent exponential service, FCFS, infinite buffer and steady state. M/M/c uses Erlang C. At ρ≥1 no finite steady-state delay is returned. K has units of years of annual baseline cost; dividing by lifetime makes P dimensionless. At κ≥1 there is no positive savings payback.')
    h('3.2 FMEA before and after mitigation')
    table(['Failure mode','Effect','S','O','D','Pre','Mitigation','Post','Evidence'],[[fmea[i-1][0] if i<3 else f'M09.S3.fmea{i}.mode',*[f'M09.S3.fmea{i}.{k}' for k in ['effect','severity','occurrence','detection','rpn_pre','mitigation','rpn_post','evidence']]] for i in range(1,6)],[1.05,.8,.3,.3,.3,.4,2.3,1.2,.75])
    p('RPN=S×O×D on 1–10 scales. Ratings are engineering judgments, not empirical probabilities. Post severity is unchanged; post occurrence/detection assumptions are shown in each product. Code and simulations demonstrate mitigations, not the numeric risk reduction; students must review these ratings.')
    fig(G/'docs/figures/capstone_economics_small.png')
    p('Scenario graphs regenerate with analytics/plot_capstone.py. The baseline point is κ=0.22, P=0.076923, NCRI=0.78, ρ=2/3, Wq=6.667 working hours. Labor ±30% gives 218.4–405.6 h/year. None of these is an observed field outcome.')
    p('Sources: MIT OCW 2.854, M/M/1 Queue (2016); NIST Two-Sample t-Test. URLs: docs/CAPSTONE_VERIFICATION_REPORT.md. Limits: ideal rays/exact pose, rigid terrain; A* + PD + analytic IK + instantaneous wrench allocation, not MPC/SLAM. No hardware, human-workload or conformity validation.')
    h('4 Empirical telemetry and independent evaluation',page=True)
    fields('M09.S4',[(x,l) for x,l in [('baseline_condition','Baseline'),('n_runs','Sample size'),('randomised','Randomization'),('log_rate_hz','Actual logging'),('seeds_recorded','Seeds recorded'),('real_csv_path','Evidence files')]])
    p('Predeclared primary outcome: terminal XY goal error, including failures; two-sided α=0.05. Both policies use the same evaluator-only path generation method, never fed to control. Cross-track RMSE uses the nearest point on that fixed polyline after 2.4 s; this is geometric tracking, not a timed trajectory error. Effort is ΣₜΣⱼτⱼ²Δt over all 16 actual actuator torques at 500 Hz, not energy.')
    table(['Measure','N','Baseline mean ± SD','Proposed mean ± SD','Unit','Direction'],[[label,*['M09.S4.'+key+'.'+k for k in ['n','baseline','proposed','unit','better']]] for key,label in [('rmse','Tracking RMSE'),('effort','Control effort'),('metric','M09.S4.metric.name')]],[1.35,.65,1.5,1.5,.9,1.5])
    fields('M09.S4',[(x,l) for x,l in [('test','Test'),('t_df_p','Primary t, df, p'),('cohens_d','Pooled Cohen d'),('ci95','95% CI')]])
    p('Strict acceptance: proposed 30/30; baseline 0/30. Every proposed run completed both 8 cm ascents and reached the goal with zero harmful contact steps; all baseline runs timed out and contacted hazards. Neither arm rolled over. The unusually large effect reflects a weak reactive baseline failing near the first barrier with low variance, not general superiority or field reliability. Mann–Whitney sensitivity p=3.02×10⁻¹¹. No failed runs were discarded.')
    fig(G/'docs/figures/capstone_results.png')
    p('Actual seed 32000 route and all 60 terminal errors. Total effort and RMSE depend on unequal run durations and failure states; lower effort alone is not efficiency. Each roughness level has 10/10 proposed acceptance and 0/10 baseline acceptance. Results from historical 100-run/26-run/7-run studies are archived separately and are not pooled here.')
    table(['Mandatory provenance declaration','Answer'],[[label,'M09.S4.'+key] for key,label in [('prov_own_runs','Technical metrics are from actual MuJoCo runs of this controller; AI help disclosed.'),('prov_no_simulated','No starter SIMULATED targets used as measured outcomes.'),('prov_plots_reproducible','Figures regenerate from retained measured CSV files.'),('prov_economics_sourced','Every business input is sourced or explicitly an assumption.')]],[6.8,.6])
    p(f'Independent raw-data check: {v["CSV_parts"]} CSV parts; continuous 2 ms samples, SHA256, effort, RMSE, contacts, goal error and both lift/support gates verified. Frozen simulator source SHA256: {stats["source_sha256"]}.')
    h('5 Showcase AI audit and peer review',page=True)
    fields('M09.S5',[(x,l) for x,l in [('linkedin_url','LinkedIn URL'),('video_length','Published duration'),('live_mujoco','Video live viewport'),('business_graph','Video business graph'),('all_members_speak','All members speak')]])
    p('Video fields are pending, so answers above are No until the actual 3–4 minute recording exists. Record live FSM, obstacles and both climbs, business sensitivity graph, and every member’s webcam/speech. Use docs/SHOWCASE_RECORDING_GUIDE.md. No slides required. Prior clips are replay evidence, not a completed showcase.')
    table(['AI case','Tool','Claim','Checked by','Correction and evidence','Error type','Verdict'],[[f'AI-0{i}',*[f'M09.S5.ai{i}.{k}' for k in ['tool','said','caught_by','truth','error_type','verdict']]] for i in range(1,4)],[.4,.5,1.3,1.25,2.15,1.1,.7])
    p('Audit cases document assistant errors and corrections; they do not assert that a named student caught them. Each student must enter their own learning and genuine contribution evidence below.')
    h('5.2 Assigned peer review')
    table(['Group','Problem','Physics','Business','Data','Clarity','Best mm:ss','Viva question'],[[name,*[f'M09.S5.peer_{key}.{k}' for k in ['problem','physics','business','data','clarity','best_moment','question']]] for key,name in [('M10','10 Surgical assistant'),('M03','03 AED delivery')]],[1.2,.65,.65,.65,.65,.65,.8,2.15])
    h('6 Individual learning record')
    p('Unmarked: members answer honestly. Role fidelity 1–5; cross-disciplinary skill Yes/No. Commit share = personal commits ÷ group commits in the assigned folder. Personal hours, experience, authorship and declarations have not been invented.')
    table(['Roll','Hours','Role 1–5','New skill','Scripts authored','Commit %','AI draft %','AI errors caught'],[[roll,*[f'M09.S6.{roll}.{k}' for k in ['hours','role_fidelity','cross_disciplinary_skill','scripts_authored','commit_share','ai_draft_share','ai_errors_caught']]] for roll in ['E064','E070','E073']],[.55,.5,.65,.65,2.2,.65,1,1.2])
    table(['Roll','Prior simulator','Most valuable learning','One course change'],[[roll,*[f'M09.S6.{roll}.{k}' for k in ['prior_simulator','most_valuable','change_one_thing']]] for roll in ['E064','E070','E073']],[.55,1.15,2.85,2.85])
    h('Individual academic sign off')
    p('Each member attests (1) original work with sources and AI use declared; (2) code/model/business scripts run; (3) whole-system and personal-role viva readiness. Optional consent permits anonymised education research, does not affect marks and may be withdrawn until results are published.')
    table(['Roll','Original','Code runs','Viva ready','Signature','Research consent'],[[roll,f'M09.SIGN.{roll}.original',f'M09.SIGN.{roll}.code_runs',f'M09.SIGN.{roll}.viva_ready','________________',f'M09.SIGN.{roll}.research_consent'] for roll in ['E064','E070','E073']],[.55,.8,.8,.8,3.1,1.35])
    p('Submit the exact same signed four-page PDF individually on Teams. Commit this DOCX and PDF in Group_09_Arnav_Vihan_Pratik/docs/ and make genuine personal-account code contributions. After entering personal fields, re-export and confirm four pages. The IEEE paper is a separate artifact; this dossier is the required form.')
    assert set(used)==set(controls),(len(used),set(controls)-set(used))
    assert len(values)==146,len(values)
    footer=sec.footer.paragraphs[0];footer.text='Group 09 | M20 technical dossier | '
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
    out=G/'docs/MDRIIA_Group_09_Technical_Dossier.docx';doc.save(out)
    final=E.fromstring(zipfile.ZipFile(out).read('word/document.xml'))
    finalcontrols={s.xpath('./w:sdtPr/w:tag/@w:val',namespaces=NS)[0]:s for s in final.xpath('//w:sdt',namespaces=NS)}
    for k,s in finalcontrols.items():
        assert E.tostring(s.find('w:sdtPr',NS),method='c14n')==original_props[k],k
        if k not in values:assert s.xpath('.//w:t/text()',namespaces=NS)==controls[k].xpath('.//w:t/text()',namespaces=NS),k
    (G/'docs/provenance/capstone_form_controls.json').write_text(json.dumps({'controls':212,'technical_fields':146,'personal_controls_preserved':66,'all_control_properties_preserved':True,'original_sha256':hashlib.sha256(SOURCE.read_bytes()).hexdigest(),'values':values},indent=2),encoding='utf-8')
    print(out)
if __name__=='__main__':main()
