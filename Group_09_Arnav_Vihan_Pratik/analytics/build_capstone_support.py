"""Build the detailed current run guide and extended technical report."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
G=Path(__file__).resolve().parents[1];A=G/'analytics';D=G/'docs'

GUIDE=[('M20 complete step by step run guide',r'''
This guide runs the Group 09 M20 quadruped project from a standalone assigned group folder. The current release uses a dense industrial yard, three alternating barriers and two 8 cm terrace climbs. It supplies live physics with FSM overlays, actual 500 Hz telemetry, an interactive business model and reproducible final-study statistics.

The robot is the supplied hybrid canine quadruped: four articulated legs with wheels at the feet, 12 leg motors and four hub motors. It rolls between obstacles and lifts front/rear leg pairs to climb. It is not a car replacement with cosmetic legs, and it is not claimed to walk continuously with every wheel braked.

Start with the live simulation and business dashboard. Use the experiment and document sections later when you need to reproduce or submit evidence. A replay is useful for inspection but is not the live simulation required for the showcase.

Quick route through this guide: pages 2–3 install and test; pages 4–5 live demo and business model; page 6 evidence/replay; page 7 full experiment; page 8 analysis and dossier; page 9 recording; page 10 GitHub and individual submission.

Confirmed team: Arnav Saurabh Surve, E064, SAP 70362400067; Vihan Shripad Joshi, E070, SAP 70362400086; Pratik Mangesh Gaikwad, E073, SAP 70362300032. The faculty IDs were confirmed by the user. Names identify the team, not verified personal code authorship.

Verified reference: 60 independent final-course runs, 30 per policy; 4,495,972 actual 500 Hz samples in 254 CSV parts. Proposed controller: 30 accepted missions, both climbs and no harmful contacts in each. Reactive baseline: 30 timeouts with harmful contacts. This simple baseline is substantially weaker than the integrated policy; do not generalize the effect size to expert teleoperation or hardware.

Files to keep: the entire Group_09_Arnav_Vihan_Pratik folder, including M20/meshes. Do not run from inside a ZIP preview. Do not copy only the Python scripts. The group folder is self-contained; a sibling M20 folder is no longer required.
'''),('1 Extract the project and prepare Windows',r'''
1. Extract the final submission ZIP to a normal local directory with enough free disk space. The full evidence archive is large because it retains actual measured CSVs and replay states. Avoid extracting into another project with identically named files.

2. Open the extracted Group_09_Arnav_Vihan_Pratik folder. Confirm that run_demo.ps1, requirements-lock.txt, M20/M20.xml, M20/meshes, src, analytics, tests and docs are present. The GitHub-only package omits videos and historical replay binaries; the final 500 Hz CSV evidence is retained.

3. Install CPython 3.12 for Windows if it is not already installed. The verified environment used Python 3.12.14 and MuJoCo 3.11.0. The pinned requirements are authoritative; do not upgrade packages during a recorded comparison.

4. Open PowerShell in the group folder. In File Explorer, type powershell in the address bar and press Enter. Check the directory and Python:
```
Get-Location
Get-ChildItem
py -3.12 --version
```
5. If the py command is unavailable, use the full path to a Python 3.12 executable. Do not use Python 3.13 merely because it is the first interpreter on PATH.

6. Run setup. This creates a local .venv and installs the exact lock file. Network access is normally needed the first time.
```
powershell -NoProfile -ExecutionPolicy Bypass -File .\run_demo.ps1 -Setup
```
An explicit interpreter is also supported:
```
powershell -NoProfile -ExecutionPolicy Bypass -File .\run_demo.ps1 -Setup -Python 'C:\path\to\Python312\python.exe'
```
Replace that example path with the real interpreter path. ExecutionPolicy Bypass applies only to this PowerShell process; no global policy change is needed. Wait for Setup complete. If pip fails, read its error, restore connectivity and rerun setup.
'''),('2 Verify the environment before demonstrating',r'''
Use the environment inside the extracted group folder. Activation is optional because every command below invokes its Python explicitly.
```
.\.venv\Scripts\python.exe --version
.\.venv\Scripts\python.exe -m pip check
.\.venv\Scripts\python.exe -m pytest -q
```
The current test suite checks model anatomy, actuator limits, mapping, ray behavior, delayed commands, acceptance gates, measured logging, statistics and business mathematics. The final verification report records the actual pass count. A test failure is a reason to investigate, not to edit a receipt manually.

Check the model directly:
```
.\.venv\Scripts\python.exe -c "import mujoco; m=mujoco.MjModel.from_xml_path('M20/M20.xml'); print(m.nq,m.nv,m.nu)"
```
Expected dimensions are 23, 22, 16. A missing STL error usually means the folder structure was broken during extraction. Restore M20/meshes and retain relative paths.

Run the default full course headlessly:
```
.\.venv\Scripts\python.exe M20\simulation.py --mode course_climb --seed 42 --out analytics\local_runs\first_course
```
Wait for the measured summary. Seed 42 without added perturbations was verified at 121.394 simulated seconds, two climbs, zero harmful contact steps, accepted=true and goal error about 0.107744 m. Wall time depends on the computer. Compare summary fields, not an animation impression.

Exit code 0 means accepted. Exit code 2 means the measured acceptance criteria were not met; this is expected for intentionally blocked or short runs. A 15-second diagnostic does not have time to finish a two-climb course and must not be described as a failed model installation.

If the display fails but headless physics works, check graphics drivers and local desktop availability. Do not replace failed rendering with a fabricated screenshot. On low-memory systems, close unrelated programs and avoid running multiple full benchmarks together.
'''),('3 Run and explain the live robot demo',r'''
Launch the native viewer from the group folder:
```
powershell -NoProfile -ExecutionPolicy Bypass -File .\run_demo.ps1 -Live
```
Equivalent direct command:
```
.\.venv\Scripts\python.exe src\live_m20_viewer.py --mode course_climb --seed 42 --out analytics\local_runs\live_demo
```
The viewport is actual stepped MuJoCo physics. Its overlay reports mode/policy, simulated time, current FSM state, controller decision, harmful contact steps, verified climb count and the actual 500 Hz saved telemetry rate. The camera follows the robot.

Watch the sequence: WAKEUP raises the body; DETOUR_NAV maps and routes around alternating barriers; DOCK_ALIGN aligns before the step; paired front and rear transfer phases raise and re-support the legs; CLEAR_REAR_EDGE moves clear of the lip; the controller repeats for the second terrace and approaches the goal.

Do not call the rise of the torso alone a verified climb. The evaluator requires separate front/rear airborne and support gates, progress across the measured lip, goal arrival and no harmful contact or rollover. At the end the overlay displays measured status and acceptance. Close the window to exit after the final summary.

Closing the window early records user_closed and does not fabricate completion. Run again into a new output folder if a complete demonstration is needed. Never overwrite a submitted trial to replace an inconvenient failure.

The course is a static industrial yard with cylinders, crates, pipes, walls, three alternating tall barriers and low terraces. Mapping uses ideal rays and exact simulator pose. The reference used for scoring is generated offline for evaluation and is not sent to the controller. Unknown terrain means the controller builds its map during motion, not that the simulator lacks geometry.

Live rendering can be slower than simulated time, especially while recording or performing disk-heavy work. Rehearse on the actual presentation computer. The --fast flag removes real-time throttling; it cannot make a slow physics/rendering computer meet a fixed wall-clock speed. Do not label an accelerated or edited segment as uninterrupted real time.
'''),('4 Run the business model beside the simulation',r'''
Open a second PowerShell window in the same group folder. Leave the live robot window running and launch the interactive business model:
```
powershell -NoProfile -ExecutionPolicy Bypass -File .\run_demo.ps1 -Business
```
Arrange the two windows side by side. The dashboard shows normalized payback, NCRI, queue delay and labor reallocation. Drag operating parity, capital parity or requests/week. Reset restores the base inputs. Explain why the curves move, rather than reading a single favorable number.

The base scenario assumes operating parity 0.22, annual baseline value 1, capital parity 0.30, five-year life, 312 reclaimed hours/year and 2080 full-time hours/year. It gives 4.615 months payback, lifetime fraction P=0.076923, NCRI=0.78 and 0.15 FTE. These are planning assumptions, not plant observations.

The queue assumes λ=8 requests/week, μ=12 services/week, one server and 40 working hours/week. Utilization is 2/3 and queue wait is 6.667 working hours. Increase requests to 12 or more: the dashboard must report no finite steady-state delay. This is the stability boundary, not a negative wait or a software failure.

Recompute machine-readable outputs and the separate telemetry-linked scenario:
```
.\.venv\Scripts\python.exe analytics\hazardous_recon_economics.py
.\.venv\Scripts\python.exe analytics\telemetry_business_bridge.py
.\.venv\Scripts\python.exe analytics\business_dashboard.py --headless
```
The bridge reads actual accepted fractions and attempt durations from the 60-run CSV. It adds assumed inspection/reset overhead of 1.5, 3 or 4.5 hours and computes hypothetical successful-job capacity. It does not convert two minutes of simulated traversal into a measured plant service rate. All assumptions and zero-success cases are retained in telemetry_business_bridge.json.

Currency-free ratios are deliberate. Increasing capital delays payback; increasing operating parity reduces NCRI; near-saturated queues grow rapidly. No NASA-TLX score, bandwidth-saving percentage or human risk reduction is inferred from robot motion.
'''),('5 Inspect telemetry and distinguish replay from live physics',r'''
Each run saves summary.json, events.json, telemetry_schema.json and numbered *_real_trials.csv parts. The scientific CSVs contain actual 2 ms samples: pose, velocities, evaluator reference, tracking error, all 16 actuator torques, FSM phase, decision, events, contact flag and support count. Parts are limited to 20,000 rows so no single raw CSV is a huge repository object.

The schema lists actuator names, units, reference polyline, sample count and each CSV SHA256. Read it before interpreting columns. The error is chassis position minus nearest point on a fixed evaluator-only path. It is not error against a time-synchronized desired motion. Integrated squared torque is effort, not electrical energy or battery consumption.

telemetry.csv and states.npz are separate 50 Hz visualization/support diagnostics. They are not the 500 Hz scientific log. states.npz permits deterministic state replay; the saved full-rate CSV is the basis for the final RMSE and effort calculations.

Replay a measured run, optionally recording a clip:
```
.\.venv\Scripts\python.exe analytics\render_yard.py analytics\local_runs\first_course --show --no-record
```
For a video, omit --no-record. Verify supported options with --help before choosing output settings. Replayed frames represent saved physical states. They do not recalculate the controller in the viewport and must be labeled replay in any evidence caption.

Inspect final-study evidence in analytics/capstone_500hz. The aggregate capstone_real_trials.csv has one row per actual trial. Each numbered trial directory retains detailed samples and summary. failure_cases.csv contains the unsuccessful baseline runs; they are part of the experiment and must remain in the submission.

The old benchmark, obstacle_course and climbing_course folders are historical studies with different protocols/source hashes. Their results are valid only for those designs. Do not mix their 50 Hz traces with the new 500 Hz study or relabel their sample sizes as current trials.
'''),('6 Reproduce the complete independent experiment',r'''
The supplied final evidence is already measured. Re-running all 60 trials is optional for ordinary demonstration, but the exact command is provided for independent reproduction. Expect meaningful compute time and several gigabytes of outputs.
```
.\.venv\Scripts\python.exe analytics\run_capstone_benchmark.py --workers 3 --out analytics\capstone_rerun
```
The script writes its design before simulations: baseline seeds 31000–31029; proposed seeds 32000–32029; 10 per policy at each roughness bound 0, 3, 6 mm. Friction, barrier positions and initial pose perturbations are seeded. The deadline per trial is 180 simulated seconds. The primary test is predeclared Welch inference on terminal XY goal error.

Use fewer workers on a lower-memory computer. Do not edit M20 source, XML or meshes during a batch; the runner compares source digests. The script refuses an existing folder without --resume. Resume accepts complete trial records with the same configuration/source and refuses ambiguous partial trials. Preserve any partial folder, then choose a new output folder if needed.

The original run finished all 60 physics trials, then exposed a NumPy Boolean serialization error while aggregating. The aggregate writer was corrected to reload the already-saved normalized JSON. All original per-trial records and CSVs were retained; no physics outcome was substituted. This repair is documented in the verification report.

Analyze the rerun and verify every raw CSV:
```
.\.venv\Scripts\python.exe analytics\analyze_capstone.py --folder analytics\capstone_rerun
```
The analyzer checks continuous 2 ms timestamps, source hashes, per-file hashes, 16 torque channels, recomputed effort and RMSE, contacts, final goal distance and both ascent gates. It writes verification.json, statistics.json and failure_cases.csv. Do not report a statistics file without checking its companion verification receipt.

Plots for the supplied study regenerate with analytics/plot_capstone.py. That script deliberately reads capstone_500hz; change its input explicitly if producing a separate rerun figure. Keep the original and rerun evidence clearly named.
'''),('7 Read the statistics and complete the dossier',r'''
The required form is docs/MDRIIA_Group_09_Technical_Dossier.docx. Its PDF is four pages. The rebuild preserves every one of the 212 original tagged Word controls, including 66 personal fields left for members. Use desktop Microsoft Word; do not flatten or replace dropdowns in Google Docs.

Technical answers use the independent final-course study. Baseline goal error is 12.3103 ± 0.0931 m; proposed is 0.1064 ± 0.0082 m. The proposed-minus-baseline difference is −12.2040 m, Welch 95% CI [−12.2388, −12.1691]. The very large pooled effect size reflects all baseline trials failing early with low variance. It is not a credible estimate of performance against an expert or a modern capable planner.

Tracking RMSE is 0.5811 ± 0.1005 m versus 0.1947 ± 0.0087 m. Total squared-torque effort is 1,128,398 ± 158,460 versus 555,625 ± 36,403 N²m²s. Duration and failure differ, so do not infer motor efficiency from this comparison.

Members must fill their own hours, learning, genuine authored scripts, commit share, peer reviews, sign-off and optional consent. Review Groups 10 and 03 using the actual videos; do not invent timestamps or scores. Enter the final LinkedIn post URL only after publishing the actual team recording.

After edits, export from Word using File → Export → Create PDF/XPS. Check that all required sections remain and the result has exactly four pages. Long personal answers can increase pagination; keep them concise while truthful. Every member submits the exact same signed PDF. The optional IEEE paper is separate and is not a substitute for this dossier.

The builder analytics/build_capstone_dossier.py regenerates technical answers from the original form and measured statistics. It will reset personal entries to original placeholders. Do not run it over the team’s signed version; first preserve the completed personal copy. Supporting details are in M20_Extended_Technical_Report.docx and CAPSTONE_VERIFICATION_REPORT.md.
'''),('8 Record the three to four minute showcase',r'''
Use docs/SHOWCASE_RECORDING_GUIDE.md for the speaking plan. Target 3:45. Rehearse before publishing. Capture a running MuJoCo viewport with visible FSM and the business graph, plus every member’s webcam and audible contribution. No slide presentation is required.

Suggested allocation: 0:00–0:25 introduction and M20 identity; 0:25–1:15 mapping and obstacle detours; 1:15–2:15 paired climbs and contact gates; 2:15–3:10 business sensitivity and queue stability; 3:10–3:45 measured comparison, AI correction, limitations and repository pointer. These are editing targets, not fabricated timestamps of a completed video.

Arnav explains mass, joints, PD control and why a climb needs lift and support gates. Vihan explains ideal rays, true pose, A*, the 500 Hz schema and comparison limits. Pratik changes a business slider, explains dimensionless payback and the queue overload boundary, and identifies which inputs are assumptions. Every student should also answer questions about the whole system.

Start screen recording before the key live actions. Show actual changing FSM text. If the laptop is slow, record a longer genuine session and edit clearly labeled live excerpts to 3–4 minutes. Do not claim an edited video is an uninterrupted real-time run. Prior MP4s are physics replays and cannot alone satisfy the live/webcam requirement.

Keep the final video out of Git. Publish natively to LinkedIn from the appropriate real account, inspect the post for playback/audio and copy its genuine URL into Section 5. The current package does not assert that posting or individual recording is complete.

Before the viva, run the two windows side by side and practice explaining a failure, a formula, a raw data row and one AI mistake. Do not memorize only favorable numbers. Bring the full offline package so network failure does not prevent the simulation, graphs, PDF or raw-evidence demonstration.
'''),('9 GitHub handoff and individual submission',r'''
The assigned destination is Group_09_Arnav_Vihan_Pratik in the course repository. The GitHub-ready ZIP contains that folder, code, models, documentation and final raw CSV evidence; it excludes MP4/NPZ and local environments. The full offline ZIP additionally retains historical evidence and replay assets. Use the correct package for the destination.

Clone the course repository or use your existing authorized checkout. Copy the prepared group folder into its assigned location and inspect changes before staging. Avoid overwriting another group’s folder. Check Git identity from each student’s own account; do not use an invented author name or divide artificial commits among members.
```
git status --short
git config user.name
git config user.email
git diff --stat
git add -- Group_09_Arnav_Vihan_Pratik
git diff --cached --stat
```
Each member must review, understand and make a genuine contribution, then commit from their own account. Use meaningful messages describing actual work. Follow repository branch/review rules and permissions. This package supplies no fabricated contribution history and has not been pushed on behalf of students.

Commit the final completed DOCX and the exact same signed PDF under docs/. Confirm no .venv, cache, MP4 or large generated binary is staged. The plain 500 Hz CSV chunks are intentional empirical evidence. Network upload of the full evidence set may take time; use Git rather than a browser upload limit workaround.

Open the remote group folder after pushing and verify the expected latest code and PDF are present. Record real commit hashes in the roster only after they exist. git shortlog -sn -- Group_09_Arnav_Vihan_Pratik reports actual contribution counts; it is not proof of effort by itself.

Every member must personally attach the same signed four-page PDF on Teams and click Turn in. Team-level upload does not complete individual submission. Keep a receipt/screenshot of each actual submission. The team must attend the live expo/viva and be prepared to run both the robot and business model. This guide does not claim those personal actions have happened.
''')]

def build_doc(pages,path):
    doc=Document();s=doc.sections[0];s.page_width=Inches(8.27);s.page_height=Inches(11.69);s.top_margin=s.bottom_margin=Inches(.6);s.left_margin=s.right_margin=Inches(.65)
    st=doc.styles['Normal'];st.font.name='Calibri';st.font.size=Pt(10.5);st.paragraph_format.space_after=Pt(7);st.paragraph_format.line_spacing=1.08
    for name,size in [('Title',22),('Heading 1',17),('Heading 2',12)]:
        st=doc.styles[name];st.font.name='Calibri';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
    for st in doc.styles:
        for b in list(st.element.iter(qn('w:pBdr'))):b.getparent().remove(b)
    md=[]
    for i,(title,body) in enumerate(pages):
        h=doc.add_paragraph(title,'Title' if i==0 else 'Heading 1');h.paragraph_format.page_break_before=i>0;md.extend(['# '+title,body])
        code=False;block=[]
        for line in body.strip().splitlines()+['']:
            if line.startswith('```'):
                if not code and block:
                    doc.add_paragraph(' '.join(block));block=[]
                if code:
                    p=doc.add_paragraph('\n'.join(block));p.paragraph_format.keep_together=True
                    for r in p.runs:r.font.name='Consolas';r.font.size=Pt(8.5)
                    block=[]
                code=not code;continue
            if code:block.append(line);continue
            if line.strip():block.append(line)
            elif block:
                text=' '.join(block);block=[]
                if text.startswith('IMAGE:'):
                    p=doc.add_paragraph();p.add_run().add_picture(str(G/text[6:]),width=Inches(6.9))
                else:doc.add_paragraph(text)
    f=s.footer.paragraphs[0];f.text='Group 09 M20 | ';fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
    doc.save(path);path.with_suffix('.md').write_text('\n\n'.join(md),encoding='utf-8')

def main():
    stats=json.loads((A/'capstone_500hz/statistics.json').read_text());verify=json.loads((A/'capstone_500hz/verification.json').read_text());bridge=json.loads((A/'telemetry_business_bridge.json').read_text());c=stats['comparisons']['goal_error_m'];b=bridge['policies']['shared'];q=b['scenarios'][1]
    report=[('M20 extended technical report',f'''
This report supports the required four-page faculty dossier. The final implementation preserves the supplied M20 quadruped morphology and adds a reproducible industrial obstacle-and-climb evaluation, actual 500 Hz telemetry and a dimensionless business model. The robot completes the tested course under ideal simulated sensing; deployment, human workload and field savings remain unvalidated.

The final experiment contains 60 independent runs, 30 per policy. Proposed map-and-transfer control accepts 30/30, each with both terrace ascents and no harmful contact steps. The reactive baseline accepts 0/30; all runs time out with harmful contacts. All outcomes are retained. No baseline or historical UGV scaffold numbers are relabeled as M20 results.

The raw audit checked {verify['total_actual_samples']:,} actual samples across {verify['CSV_parts']} CSV parts. Primary final goal error differs by {c['difference_proposed_minus_baseline']:.6f} m (proposed minus baseline), Welch t={c['t']:.6f}, df={c['df']:.6f}, p={c['p_two_sided']:.8g}, 95% CI [{c['CI95'][0]:.6f}, {c['CI95'][1]:.6f}]. The pooled Cohen d of {c['cohen_d_pooled']:.6f} is very large because the weak baseline fails near the first obstacle with low variability. This is integration evidence, not proof of state-of-the-art superiority.

Canonical source is inside Group_09_Arnav_Vihan_Pratik/M20. The frozen simulation digest is {stats['source_sha256']}. The evaluation folder includes the predeclared design, source snapshot, mesh hashes, per-run configurations, measured samples, summary records and independently recomputed statistics.

AI assistance was substantial in implementation, experiments, analysis and drafting. The original mechanical model and mesh assets were reused. Team identities are confirmed, but individual authorship, hours, learning, peer scores and signatures must be personally attested. This report does not create contribution evidence.

IMAGE:docs/figures/course_climbing_frame.png

Historical rendered physical state showing paired-leg transfer. The final study’s measured records are identified by their separate source digest and trial IDs.
'''),('Mechanical model and control architecture',r'''
The M20 MJCF has a floating base, twelve articulated leg joints and four actuated foot wheels: nq=23, nv=22, nu=16. The supplied total mass is 34.525856 kg and seventeen STL meshes retain the robot’s appearance. Collision primitives and inertias are supplied/simplified model quantities, not identified hardware parameters.

MuJoCo integrates at 0.002 s with implicitfast. The robot uses friction coefficients 1/0.01/0.001; floor sliding friction is seeded U(0.75,1.15), with torsional 0.005 and rolling 0.0001. Robot solref is 0.005/1; floor solref is 0.02/1. These settings provide numerical contact compliance and are not calibrated material measurements.

Multibody dynamics obey M(q)q̈+C(q,q̇)q̇+g(q)=τ+Jᵀf. Joint commands combine PD tracking, bias compensation and stance-wrench allocation before saturation. Leg Kp is [500,800,800] N m/rad and Kd is [25,40,40] N m s/rad. Leg torque caps are 76.4 N m; hub caps are 21.6 N m. The instantaneous regularized wrench solution is not model-predictive control.

Analytic two-link inverse kinematics uses two 0.25 m sagittal links and the supplied lateral geometry. Wheel target conversion retains a 0.09 m servo radius versus 0.093 m collision radius. Effective steering width is 0.55 m. This mismatch and skid-steer approximation limit interpretation of the kinematic slip proxy; it is not a measured tire-slip estimator.

The FSM starts in WAKEUP, cruises through mapped detours, aligns for a terrace, shifts support, transfers the front pair, raises/advances the body, shifts support to the front and transfers the rear pair. A clearance phase moves the hind feet beyond the lip before the next navigation/transfer phase. A second measured terrace invokes the sequence again.

Strict climb acceptance requires front and rear pair airborne events and re-support events plus geometric progression. Mission acceptance additionally requires goal arrival, no rollover, no tilt-stop violation and zero harmful contact steps. Expected support on traversable terrain is distinguished from contacts with harmful obstacle geometry. This prevents torso rise, foot sliding or reaching a goal through collisions from being counted as successful climbing.

Worked static check: mg=34.525856×9.81=338.698647 N. Equal four-foot loading would be 84.674662 N per foot; a 0.25 m lever gives 21.168666 N m. These are illustrative equilibrium calculations. The actual seed-42 support normals near 2.302 s sum to about 338.739 N; individual normals are read from the measured diagnostic CSV and need not be equal.
'''),('Perception terrain and experimental design',r'''
The industrial yard includes three alternating tall barriers, cylinders, crates, pipe/boundary geometry and two 8 cm terrace rises (16 cm cumulative elevation). Terrain roughness adds physical low rubble boxes rather than changing labels alone. The tested environment is static, rigid and deliberately structured; it is not a loose-debris disaster simulator.

Ideal ray intersections and exact MuJoCo pose update a 2.5D elevation/traversability grid. The dense course uses 5 cm cells. Unknown cells are blocked for path search; local sensing/frontier behavior reveals reachable space. A* routing uses obstacle inflation and height/roughness/slope costs. Height bands identify the low step edge and docking point. The broad terraces are approximately aligned with world x, which is a task assumption rather than general stair detection.

The evaluator separately builds a fixed reference polyline from the known fixture geometry. Both policies are scored by the same method. This reference is never sent to control. It avoids the optimistic error obtained when an online replan begins at the current robot position, but it measures nearest-segment cross-track distance rather than time-indexed tracking performance.

Thirty independent seeds per policy were predeclared. Baseline 31000–31029 and proposed 32000–32029 do not overlap. Each arm has ten trials with roughness bound 0, ten at 3 mm and ten at 6 mm. For nonzero bounds, rubble heights are U(bound/2,bound). Floor friction is U(0.75,1.15), barriers shift within ±0.08 m, initial lateral displacement within ±0.02 m and yaw within ±0.02 rad. Delay is fixed at the assumed 0.35 s. The per-trial time budget is 180 simulated seconds.

The baseline uses the same M20 and standing/rolling actuation but a delayed reactive driver without mapping or paired climb logic. The proposed policy bundles multiple improvements. This comparison tests system integration; it cannot isolate the contribution of mapping, climb transfer or delay handling. A stronger baseline and controlled ablations are important future work.

Historical evidence remains separate: 100 original paired runs, 26 avoidance/blocked runs and seven earlier combined-course checks. Their source hashes and protocols differ. They are not pooled with the final sixty independent runs or used to inflate sample size. Earlier 50 Hz diagnostics are explicitly distinguished from the final actual 500 Hz scientific logger.
'''),('Telemetry statistics and failure interpretation',f'''
TrialLogger streams actual post-step pose and velocity with the actuator force applied over the preceding 2 ms interval. Unit actuator gear is asserted, so the sixteen force channels represent motor torque. The logger writes bounded plain CSV parts, not interpolated 500 Hz rows generated from 50 Hz data. Numeric serialization preserves timestamps at six decimals and measured channels at nine significant digits.

Tracking RMSE is sqrt(mean(error_x²+error_y²)) after the 2.4 s startup interval. Effort is the complete time integral of the sum of sixteen squared torques. Each includes failed trials and their actual durations. Separate 50 Hz state arrays support playback and diagnostic foot-force inspection.

The verifier checks all CSV hashes and continuous timestamps, run IDs/seeds, torque finiteness, position-reference arithmetic, RMSE, integrated effort, harmful contact counts, final distance to goal and both airborne/re-support gate pairs for accepted runs. The default seed-42 migration was bitwise identical in saved qpos/qvel/ctrl/targets/time to the previous controller version: adding full-rate logging did not change its dynamics.

Primary goal-error means are 12.310317 m baseline and 0.106357 m proposed. Cross-track RMSE means are 0.581105 and 0.194681 m. Squared-torque effort means are 1,128,398.226 and 555,625.351 N²m²s. Standard deviations, inference and secondary metrics are in statistics.json; all failures are in failure_cases.csv.

Welch inference was chosen for independent groups with unequal variances. The pooled-SD standardized mean difference is reported separately from the Welch standard error. A Mann–Whitney sensitivity test gives p≈3.0199×10⁻¹¹. The extreme separation is expected when one policy cannot progress through the fixture; a small p-value does not cure weak-baseline design or prove hardware safety.

IMAGE:docs/figures/capstone_results.png

The first plot uses measured seed 32000 and its fixed evaluator path. The second includes all sixty runs. Every proposed roughness stratum has 10/10 acceptance; baseline strata have 0/10. No trial was removed to improve results. These sample counts do not establish a population success probability of one.
'''),('Business assumptions and empirical telemetry bridge',f'''
The base economic model uses normalized annual baseline value V=1, annual operating cost O=0.22, capital K=0.30 in annual-baseline cost units and life L=5 years. κ=O/V; NCRI=1−κ; payback months=12K/(V−O); lifetime-normalized P=K/[L(V−O)]. These yield κ=0.22, NCRI=0.78, T=4.615385 months and P=0.076923. Scaling K,V,O by the same factor leaves the result unchanged.

Assumed reclaimed annual hours H=312 and full-time capacity 2080 give ΔFTE=0.15. A ±30% hours scenario gives 0.105–0.195 FTE. At κ≥1 there is no positive savings payback; with K=0.30 and L=5 the within-life boundary is κ=0.94. These are explicit assumptions, not measured employment changes.

M/M/1 uses λ=8/week and μ=12/week, with 40 working hours/week. ρ=2/3, Wq=λ/[μ(μ−λ)]=1/6 week=6.666667 working hours, W=0.25 week, Lq=1.333333 and L=2. M/M/c uses Erlang C; zero arrivals and unstable ρ≥1 are tested. Poisson arrivals, exponential independent service, FCFS and steady state are unverified modeling assumptions.

The separate telemetry bridge reads real final-study acceptance and attempt duration. Proposed mean attempt duration is {b['measured_mean_attempt_seconds']:.6f} simulated seconds and empirical accepted fraction is {b['measured_acceptance_fraction']:.3f}. With an explicitly assumed 3-hour inspection/reset overhead per attempt and 40 working hours/week, μ_eff=40p/(3+t/3600)={q['effective_success_capacity_per_week']:.6f} successful jobs/week. At λ=8 this scenario gives Wq={q['queue']['Wq_work_hours']:.6f} working hours. This is a sensitivity scenario, not a measured service rate; baseline zero accepted service is reported as zero capacity, not divided by zero.

IMAGE:docs/figures/capstone_economics_small.png

FMEA ratings are subjective 1–10 planning scores. Severity remains fixed after mitigation; assumed occurrence/detection reductions produce the displayed post RPNs. Simulation/code supports tilt latches, local stopping, blocked-path handling, contact gates and witness validation. It does not statistically validate the numerical risk scores or establish regulatory conformity.
'''),('Reproduction audit and remaining human work',r'''
Run the group-only commands in the complete guide. The simulation module and assets are now under the assigned group folder, and models use relative ../M20/meshes references. A clean group-only staging check runs tests and a full seed-42 course without a sibling M20 dependency. The final verification report records the resulting receipt rather than assuming portability from a successful local run.

The original final batch completed all sixty simulations before an aggregate JSON serialization bug surfaced on a NumPy Boolean. The batch worker now returns the normalized per-run JSON; the complete evidence was re-aggregated and independently audited. No raw sample or outcome was regenerated to hide a failure. This is separate from the physical source digest, which remained frozen throughout the experiment.

Live UI evidence confirms the native viewport displays FSM/decision/contact/climb telemetry. The business dashboard was opened and its sliders changed: capital altered payback, and requests above capacity produced the unstable-queue message. Genuine UI screenshots are distinguished from renderer-generated physics replays. A short window-closure diagnostic records user_closed and is not claimed as a complete course.

The four-page faculty form retains 212 tagged controls and all original dropdown properties. Its 146 technical fields are completed; 66 personal controls remain unchanged for students. Word-to-PDF output is rendered and visually inspected. The extended report supplies detail without making the required dossier eighteen pages long. The existing unverified dossier target is updated after final QA, while its original backup is retained.

Sources for statistical formulas: NIST/SEMATECH e-Handbook, Two-Sample t-Test for Equal Means, https://www.itl.nist.gov/div898/handbook/eda/section3/eda353.htm. Queue formulas: MIT OpenCourseWare, 2.854 Introduction to Manufacturing Systems, M/M/1 Queue (2016), linked in CAPSTONE_VERIFICATION_REPORT.md. Robotics literature metadata and DOI checks are retained in docs/bibliography/verification.json. No field or human outcome is attributed to those sources as a result of this project.

Remaining human actions are deliberate: genuine personal-account contributions, student review/understanding, webcam/speech recording, native LinkedIn posting and real URL, assigned peer reviews, personal learning fields and signatures, optional consent choices, and each member’s individual Teams Turn in. These cannot be completed by manufacturing records. The code, mathematical scenarios and measured evidence are prepared for those actions.

Future evaluation should use held-out geometry, stronger comparable-capability baselines, perception noise and drift, dynamic hazards, more severe roughness, hardware identification, energy measurement and actual operational workload/time studies. The current evidence establishes the tested simulated integration and nothing beyond it.
''')]
    build_doc(GUIDE,D/'M20_Complete_Step_by_Step_Run_Guide.docx')
    build_doc(report,D/'M20_Extended_Technical_Report.docx')
    print('Built 10 guide sections and 6 extended-report sections; render/page QA required.')
if __name__=='__main__':main()
