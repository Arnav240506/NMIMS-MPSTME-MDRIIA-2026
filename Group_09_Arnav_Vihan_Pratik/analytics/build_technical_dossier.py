"""Fill the supplied dossier form without removing its tagged Word controls.

Run build_dossier_evidence.py first. Optional document dependencies are listed
in requirements-docs.txt. Personal declarations are deliberately preserved.
"""
from pathlib import Path
from copy import deepcopy
from datetime import datetime,timezone
import json,hashlib,zipfile,argparse
from lxml import etree as E
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT=Path(__file__).resolve().parents[2];GROUP=ROOT/'Group_09_Arnav_Vihan_Pratik';HERE=GROUP/'analytics'
NS={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'};W=NS['w'];R='http://schemas.openxmlformats.org/package/2006/relationships'
INPUT_HASH='a920ee6d76a9981ca40a2d326e88923ba34cceefd2b25a201753d9ed96a2de53'

def tag(s):return s.find('w:sdtPr/w:tag',NS).get(qn('w:val'))
def canon(e):return E.tostring(e,method='c14n',exclusive=True)
def protected(t):
    return t.startswith(('M09.S6.','M09.SIGN.','M09.S5.peer_')) or (t.startswith('M09.S2.script') and t.endswith(('.author','.authorship')))
def text_of(e):return ''.join(e.xpath('.//w:t/text()',namespaces=NS))
def run_text(text,size=8,font='Calibri',bold=False):
    r=E.Element(qn('w:r'));rp=E.SubElement(r,qn('w:rPr'));E.SubElement(rp,qn('w:rFonts'),{qn('w:ascii'):font,qn('w:hAnsi'):font})
    E.SubElement(rp,qn('w:sz'),{qn('w:val'):str(int(2*size))});E.SubElement(rp,qn('w:color'),{qn('w:val'):'000000'})
    if bold:E.SubElement(rp,qn('w:b'))
    for i,line in enumerate(str(text).split('\n')):
        if i:E.SubElement(r,qn('w:br'))
        t=E.SubElement(r,qn('w:t'));t.set('{http://www.w3.org/XML/1998/namespace}space','preserve');t.text=line
    return r
def paragraph(text,size=9,bold=False):
    p=E.Element(qn('w:p'));pr=E.SubElement(p,qn('w:pPr'));E.SubElement(pr,qn('w:spacing'),{qn('w:after'):'50'})
    p.append(run_text(text,size,bold=bold));return p

def build(source,output):
    assert hashlib.sha256(source.read_bytes()).hexdigest()==INPUT_HASH,'Input form changed; review its slots before rebuilding'
    with zipfile.ZipFile(source) as z:raw={n:z.read(n) for n in z.namelist()}
    root=E.fromstring(raw['word/document.xml']);body=root.find('w:body',NS);tables=body.findall('w:tbl',NS);paras=body.findall('w:p',NS)
    controls={tag(s):s for s in root.findall('.//w:sdt',NS)};assert len(controls)==212 and len(tables)==29
    protected_before={t:canon(s) for t,s in controls.items() if protected(t)}
    props_before={t:canon(s.find('w:sdtPr',NS)) for t,s in controls.items()}
    roster=json.loads((GROUP/'docs/TEAM_ROSTER.json').read_text());stats=json.loads((HERE/'benchmark/statistics.json').read_text());metrics=json.loads((HERE/'dossier/metrics.json').read_text())
    course=json.loads((HERE/'climbing_course/statistics.json').read_text());avoid=json.loads((HERE/'obstacle_course/statistics.json').read_text());hero=json.loads((HERE/'climbing_course/00_course_climb_42_shared/summary.json').read_text())
    filled={}
    def fill(t,value):
        assert t in controls and not protected(t),t
        s=controls[t];pr=s.find('w:sdtPr',NS);options=pr.xpath('./w:dropDownList/w:listItem/@w:displayText',namespaces=NS)
        if options:assert str(value) in options,(t,value,options)
        ts=s.xpath('./w:sdtContent//w:t',namespaces=NS);assert ts
        ts[0].text=str(value);ts[0].set('{http://www.w3.org/XML/1998/namespace}space','preserve')
        for extra in ts[1:]:extra.text=''
        for run in s.xpath('./w:sdtContent//w:r',namespaces=NS):
            rp=run.find('w:rPr',NS)
            if rp is None:rp=E.Element(qn('w:rPr'));run.insert(0,rp)
            for key,val in [('sz','16'),('szCs','16'),('color','000000')]:
                node=rp.find('w:'+key,NS)
                if node is None:node=E.SubElement(rp,qn('w:'+key))
                node.set(qn('w:val'),val)
        filled[t]=str(value)
    def cell(i,j,k):return tables[i].findall('w:tr',NS)[j].findall('w:tc',NS)[k]
    def plain(i,j,k,value,size=8,font='Calibri'):
        c=cell(i,j,k);assert not c.findall('.//w:sdt',NS),(i,j,k)
        p=c.find('w:p',NS)
        if p is None:p=E.SubElement(c,qn('w:p'))
        for old in list(p):
            if old.tag!=qn('w:pPr'):p.remove(old)
        p.append(run_text(value,size,font))
        for extra in c.findall('w:p',NS)[1:]:c.remove(extra)
    def ptext(i,value):
        p=paras[i]
        for old in list(p):
            if old.tag!=qn('w:pPr'):p.remove(old)
        p.append(run_text(value,9))
    def after(table,text):table.addnext(paragraph(text,8))

    # Correct the existing form, keeping its section order and administrative fields.
    ptext(2,'GROUP 09 M20 TECHNICAL DOSSIER')
    pr=paras[2].find('w:pPr',NS)
    if pr is None:pr=E.Element(qn('w:pPr'));paras[2].insert(0,pr)
    st=pr.find('w:pStyle',NS)
    if st is None:st=E.SubElement(pr,qn('w:pStyle'))
    st.set(qn('w:val'),'Title');paras[2].remove(paras[2].find('w:r',NS));paras[2].append(run_text('GROUP 09 M20 TECHNICAL DOSSIER',14,bold=True))
    plain(0,0,1,'Group 09 | M20 hybrid canine quadruped for unknown hazardous-terrain reconnaissance')
    for j,m in enumerate(roster['members'],1):plain(1,j,2,m['sap_id']);plain(1,j,4,m['viva_focus'])
    ptext(5,'Identifiers follow the approved project Master Prompt and current TEAM_ROSTER.json. Roles are assigned review and viva responsibilities, not evidence of individual code authorship. Personal declarations, peer scores, signatures and optional consent remain for the members.')
    ptext(18,'Audit resolution: unsupported questionnaire targets are excluded from the analysis. No NASA-TLX respondents or scores were collected. Economics uses explicitly illustrative parities; no measured human workload or field-safety benefit is claimed.')
    ptext(31,'Keep economic quantities dimensionless and label each uncalibrated input as an assumption. Vary uncertain inputs and report the break-even condition; simulation acceptance does not establish observed operational savings.')
    after(tables[0],'Technical revision 26 October 2026. This completes the supplied dossier form and appends an engineering annex. The form states a 20-mark capstone structure; the repository also supplies a 100-mark ICA rubric. Both are mapped in Annex A9 without claiming an awarded grade.')
    for i in [1,2]:
        fill(f'M09.S1.claim{i}.source','No primary evidence measured or supplied; exclude this numeric outcome. No bandwidth experiment or human workload study exists in this project.')
        fill(f'M09.S1.claim{i}.source_type','No source found');fill(f'M09.S1.claim{i}.verdict','Delete')
    fill('M09.S1.claim3.source','Design assumption, not a measured network statistic: 0.35 s default; original trials cycle 0.05, 0.35, 0.8 and 1.2 s. Source: TrialConfig and LatencyQueue.')
    fill('M09.S1.claim3.source_type','No source found');fill('M09.S1.claim3.verdict','Correct')
    standards={
        'iso10218':('No','Not used as a governing conformity target: no industrial manipulator integration is assessed.'),
        'iso3691':('Partly','Safety concepts only; ISO 3691-4:2023 scope/Annex A reviewed. No conformity claim. The 1.5 m slowdown is implemented; 0.20 m stop differs from the course 0.40 m value. No personnel detection or certified brake test.'),
        'ria':('No','No industrial arm or collaborative manipulator application is evaluated; no compliance claim.'),
        'iec60601':('No','The model is not medical electrical equipment and has no medical-device evaluation.'),
        'domain':('Partly','SI and x-forward/y-left/z-up conventions are used. No running ROS frame tree or SLAM. Explosive, military and severe-environment hazards are outside the tested rigid yard.')}
    for key,(applies,why) in standards.items():fill('M09.S1.std_'+key+'.applies',applies);fill('M09.S1.std_'+key+'.clause',why)
    fill('M09.S1.baseline_value','Original 50 paired scenes: delayed software baseline accepted 10/50 (20%); map/transfer policy accepted 47/50 (94%). Final goal error: 1.358 +/- 1.161 m vs 0.218 +/- 0.536 m.')
    fill('M09.S1.baseline_source','Actual MuJoCo benchmark/results.csv and benchmark/statistics.json; all 100 raw trials retained. This is a software-driver baseline, not a field or human baseline.')
    fill('M09.S1.stakeholder','Potential inspection/response personnel and supervisors. Their present exposure, labor cost and workload were not measured; only the robot simulation and assumed economic scenario were evaluated.')
    plain(6,0,1,'Recovered M20 hybrid quadruped: 12 leg joints, four foot-wheel joints, 17 original STL meshes, 34.525856 kg total supplied mass; nq=23, nv=22, nu=16.')
    plain(6,1,1,'M20/M20.xml (original anatomy); M20/M20_obstacle_course.xml and models/m20_obstacle_course.xml (current course). The legacy hazardous_terrain_ugv.xml filename now resolves to the M20 yard.')
    plain(6,2,1,'M20/recon_controller.py, hybrid_control.py and simulation.py; src/rough_terrain_recon_controller.py is a compatibility entry point.')
    plain(6,3,1,'2 ms implicitfast integration, pyramidal contacts. Robot friction 1 / 0.01 / 0.001; floor sliding friction U(0.75,1.15), torsional 0.005, rolling 0.0001. Robot solref 0.005 1; floor default 0.02 1. Four leg chains plus four hub motors; all 16 torques bounded.')
    plain(6,4,1,'WAKEUP -> mapped detour -> DOCK_ALIGN -> APPROACH -> PACK_REAR_STANCE -> SHIFT_TO_REAR -> FRONT_PAIR_STEP -> TORSO_RISE -> BODY_ADVANCE -> SHIFT_TO_FRONT -> REAR_PAIR_STEP -> SETTLE -> CLEAR_REAR_EDGE -> PLATEAU_RECON. Repeat for a higher measured terrace; tilted/no-route/timeout cases remain failures.')
    plain(6,5,1,'Multibody dynamics are integrated by MuJoCo. Command law: tau = sat[Kp(q_ref-q)-Kd*qdot+qfrc_bias+tau_stance]. Instantaneous regularized wrench allocation is not MPC. Analytic two-link leg IK is used. Effective steering B=0.55 m and hub servo radius=0.09 m; collision radius=0.093 m. Base risk c=0.4*slope_norm+0.3*roughness_norm+0.3*step_norm, with a separate obstacle-proximity penalty.')
    fill('M09.S2.timestep','0.002 s (500 Hz physics and torque updates); 50 Hz state/CSV samples; 100 Hz wrench allocation, 10 Hz grid cost, 5 Hz planner requests.')
    fill('M09.S2.friction','Robot geoms 1 / 0.01 / 0.001; floor sliding U(0.75,1.15), with 0.005 / 0.0001. Seed 42 floor coefficient 1.05958242. Rigid Coulomb-style contact; not a calibrated tire model.')
    fill('M09.S2.solref','Robot 0.005 s / 1; floor 0.02 s / 1. At the evaluated settled contact, mixed solref is 0.0125 / 1. These are numerical settings, not measured material parameters; solimp defaults 0.9 / 0.95 / 0.001 / 0.5 / 2.')
    fill('M09.S2.control_law','Other')
    fill('M09.S2.gains','Leg joint PD Kp=[500,800,800] N m/rad; Kd=[25,40,40] N m s/rad. Wheel roll gain 3 N m s/rad; brake Kp=60 and Kd=5. Leg limits 76.4 N m, hubs 21.6 N m. Wrench regularizer 1e-4. Analytic IK plus instantaneous stance wrench; no integral term or predictive horizon.')
    fill('M09.S2.fsm_states','WAKEUP, DETOUR_NAV/CRUISE_MAP, DOCK_ALIGN; paired sequence listed above; CLEAR_REAR_EDGE, PLATEAU_RECON. Reasons include no_safe_path, obstacle_slow, obstacle_stop_replan and tilt_stop. See Annex A2 for transitions and gates.')
    fill('M09.S2.model_loads','Yes')
    changes=['Full M20 model alias and physical yard, retaining original anatomy/mesh provenance.','Compatibility entry point to the measured M20 runner; detailed control lives in M20 modules.','Dimensionless assumed parity, payback and FTE equations; no questionnaire target scores.','Measured ideal-ray grid, unknown blocking, A*, height-banded step-edge measurements.','Time-based command queue, mapped navigation, contact-gated repeated terrace transfer and strict acceptance.']
    for i,x in enumerate(changes,1):fill(f'M09.S2.script{i}.change',x)
    fill('M09.S2.script4.name','M20/m20_rosbot_mapping.py');fill('M09.S2.script5.name','M20/recon_controller.py and simulation.py')
    after(tables[8],'Authorship fields remain unselected because no student-specific commit evidence is available. The supplied model was reused and implementation/experiments/documentation received substantial AI assistance. Assigned roles do not establish who wrote a file; each student must review and personally attest their own contribution.')
    code=(ROOT/'M20/recon_controller.py').read_text().splitlines();start=next(i for i,l in enumerate(code) if l.startswith('def track('));end=next(i for i,l in enumerate(code[start+1:],start+1) if l.startswith('def reactive_driver('));snippet='M20/recon_controller.py, lines '+str(start+1)+'-'+str(end)+'\n'+'\n'.join(f'{i+1:03d} {code[i]}' for i in range(start,end) if code[i].strip());assert len(snippet.splitlines())<=25
    plain(9,0,0,snippet,8,'Consolas')
    code_pr=cell(9,0,0).find('w:p/w:pPr',NS)
    if code_pr is None:
        code_pr=E.Element(qn('w:pPr'));cell(9,0,0).find('w:p',NS).insert(0,code_pr)
    code_jc=code_pr.find('w:jc',NS)
    if code_jc is None:code_jc=E.SubElement(code_pr,qn('w:jc'))
    code_jc.set(qn('w:val'),'left')
    plain(11,1,1,'Potential industrial inspection or reconnaissance use is motivation only. The implementation demonstrates local mapping, detours and low-terrace transfer in a rigid simulated yard. No measured reduction in human exposure, radio bandwidth or cognitive workload is established.')
    plain(11,2,1,'Illustrative current inputs: operating parity 0.22; reclaimed-value parity 1.00; capital parity 0.30; assumed reclaimed task hours 312/year; reference work capacity 2080 h/year. No NASA-TLX values are used.')
    plain(12,1,1,'H_reclaimed is an assumed task-time scenario; delta_FTE = H_reclaimed / 2080.')
    plain(12,2,1,'T_months = 12*K_capex/(1-kappa), for kappa < 1. No service-life fraction is computed without an assumed lifetime.')
    plain(12,3,1,'kappa = operating parity / reclaimed-value parity; net parity margin = 1-kappa under a unit value baseline. Negative when kappa > 1.')
    plain(12,4,1,'No calibrated queue is implemented. M/M/1 would require justified arrival/service rates and a stability check; no rates are fabricated.')
    econ={'hreclaimed_inputs':'Assumption: H=312 h/year; full-time reference=2080 h/year. No field time study.', 'hreclaimed_metric':'312 h/year assumed; delta_FTE=0.150 (dimensionless).',
        'payback_inputs':'Assumptions: operating parity 0.22, reclaimed-value parity 1, capital parity 0.30.', 'payback_metric':'kappa=0.220; illustrative payback=4.615 months; not observed savings.',
        'ncri_inputs':'Assume unit baseline value parity; margin=1-kappa.', 'ncri_metric':'0.780 dimensionless assumed margin; no empirical saving claim.',
        'queue_inputs':'Not calibrated: no observed arrival rate, service rate or queue distribution.', 'queue_metric':'Not estimated. Labor-reallocation graph supplied instead.',
        'sensitivity_inputs':'Scenario kappa=0.12,0.22,0.45,0.80,1.00,1.20 at capital parity 0.30.', 'sensitivity_metric':'Payback=4.091,4.615,6.545,18.000 months for kappa<1. No finite positive payback at kappa>=1.'}
    for k,v in econ.items():fill('M09.S3.E073.'+k,v)
    fill('M09.S3.E073.queue_model','Other (explain in 3.1)');fill('M09.S3.E073.queue_params','Not estimated; no arrival/service observations or server count assumed. No queuing benefit reported.')
    fill('M09.S3.E073.payback_months','4.615384615 months, from the explicitly assumed parity inputs.');fill('M09.S3.E073.reproducible','Yes')
    fmea=[('High-centring / tip-over on rubble step','Lost mission; possible model overturn',8,5,6,3,3,'Tilt latch, torque caps, staged transfer; long walking stability unproven.'),
        ('Communication delay / late stop','Contact or mission failure',8,6,6,3,2,'Local map routing and 50 Hz ranging override delayed reactive motion.'),
        ('Unobserved / sealed passage','No feasible route to goal',6,5,6,3,2,'Unknown cells blocked; frontier checks; closed-gate failures retained.'),
        ('Step lip bias / rear-foot scrape','False completion or leg contact',7,6,5,3,2,'Measured edge coordinates, 0.24 m rear landing, separate lift/contact gates.'),
        ('Inconsistent distance diagnostic','Misleading clearance claim',5,4,7,3,2,'Check distance witness; withhold inconsistent minima; use solver contact audit.')]
    for i,(mode,effect,S,O,D,O2,D2,mitigation) in enumerate(fmea,1):
        base=f'M09.S3.fmea{i}.'
        if i>2:fill(base+'mode',mode)
        for key,value in {'effect':effect,'severity':str(S),'occurrence':str(O),'detection':str(D),'rpn_pre':str(S*O*D),'mitigation':mitigation,'rpn_post':f'{S*O2*D2} ({S}x{O2}x{D2})','evidence':'Both'}.items():fill(base+key,value)
    after(tables[14],'FMEA S/O/D values are subjective engineering planning assumptions on the supplied 1-10 scale, not estimated event probabilities or certified risk reductions. Severity is held fixed after mitigation. Post O/D ratings need student review; code/tests support the mechanisms, not the numerical ratings. The five post triples are 8x3x3, 8x3x2, 6x3x2, 7x3x2 and 5x3x2.')
    plain(16,0,1,'Original main study: 50 matched environments / 100 simulations. Separate avoidance extension: 12 matched environments plus two closed gates / 26 simulations. Separate two-climb check: six shared seeds and one illustrative baseline / seven simulations. Do not pool these designs.')
    plain(16,1,1,'Actual physics/control and contact audit at 500 Hz. Saved telemetry and full-state snapshots at 50 Hz (20 ms). Current course additionally saves time-stamped 5 Hz planner histories. The saved CSV is not a 500 Hz log.')
    plain(16,2,1,'Original policies each have 50 trials. Seeded friction, barriers, step heights and low rubble vary; initial crouch is controlled. New dense scenes vary friction and barricade offset. Obstacles are static; no randomized pedestrian/obstacle velocity study.')
    plain(16,3,1,'Primary strict mission acceptance; final goal error, contacts, time, path length, attitude, slip proxy and mechanical-work summary. Added dossier diagnostics: 50 Hz sampled torque-squared integral and six current-course rolling-path navigation RMSE values. Original full planner histories are absent.')
    plain(16,4,1,'Matched design: primary two-sided exact McNemar acceptance test and paired bootstrap interval. Final goal error has exploratory paired t/d_z and Wilcoxon sensitivity analysis because differences are non-normal.')
    plain(16,5,1,'analytics/benchmark/*/telemetry.csv + states.npz + summary.json; results.csv/statistics.json aggregate actual runs. Extensions use obstacle_course/ and climbing_course/. Dossier derived metrics: analytics/dossier/dossier_real_trials.csv and metrics.json.')
    values={'baseline_condition':'Delayed deterministic reactive software driver, same complete M20 and standing torque servo. No human operator, no measured NASA-TLX. Proposed policy bundles map planning and paired transfer; it is not a single-feature ablation.',
        'n_runs':'Original: 50 per policy (100 actual simulations). Avoidance extension: 12 per policy + two blocked fixtures. Combined: six shared and one illustrative baseline; descriptive only.',
        'randomised':'Original seeds 1000-1049; friction U(0.75,1.15), step U(0.055,0.085) m where used, barrier position/width and low rubble. Delays cycle 0.05/0.35/0.8/1.2 s. Combined course uses two fixed 8 cm rises and seeded friction/barrier offsets.',
        'log_rate_hz':'50 Hz saved state/CSV; 500 Hz physical contact audit; simulation time independent of wall-clock compute/rendering.',
        'real_csv_path':'analytics/benchmark/results.csv and each telemetry.csv; obstacle_course/results.csv; climbing_course/results.csv; dossier/dossier_real_trials.csv. Source version hashes and raw states are retained.'}
    for k,v in values.items():fill('M09.S4.'+k,v)
    fill('M09.S4.seeds_recorded','Yes')
    tr=metrics['navigation_tracking_summary'];eff=metrics['sampled_effort']
    fill('M09.S4.rmse.n','6 proposed; 0 comparable baseline');fill('M09.S4.rmse.baseline','Not available: no comparable online reference path.');fill('M09.S4.rmse.proposed',f"{tr['mean']:.5f} +/- {tr['sd']:.5f}; current navigation only")
    fill('M09.S4.rmse.unit','m');fill('M09.S4.rmse.better','Lower is better')
    fill('M09.S4.effort.n','50 per policy');fill('M09.S4.effort.baseline',f"{eff['baseline']['mean']:.1f} +/- {eff['baseline']['sd']:.1f}");fill('M09.S4.effort.proposed',f"{eff['shared']['mean']:.1f} +/- {eff['shared']['sd']:.1f}");fill('M09.S4.effort.unit','(N m)^2 s; sampled proxy');fill('M09.S4.effort.better','Lower is better')
    fill('M09.S4.metric.name','Final goal error (original paired study)');fill('M09.S4.metric.n','50 per policy');fill('M09.S4.metric.baseline','1.3575 +/- 1.1607');fill('M09.S4.metric.proposed','0.2181 +/- 0.5364');fill('M09.S4.metric.unit','m');fill('M09.S4.metric.better','Lower is better')
    after(tables[18],'RMSE above is a new six-seed rolling-path diagnostic, not the original 50-pair outcome and not a baseline comparison. Replanning resets the reference start near the robot, making this an optimistic cross-track measure. Effort integrates saved 50 Hz torque commands, not all 500 Hz commands; it is not energy. Durations and failure rates differ, so lower unnormalized effort alone does not establish efficiency.')
    plain(19,0,0,'Continuous sensitivity test; primary binary test in Annex A4')
    fill('M09.S4.test','Paired t-test');fill('M09.S4.t_df_p','Final goal error sensitivity: t=-6.761397, df=49, p=1.5540753e-8. Primary accepted-mission McNemar p=1.4551915e-11; 37 shared-only pairs, zero baseline-only.')
    fill('M09.S4.cohens_d','Paired d_z=-0.956206 for shared-minus-baseline goal error. Shapiro p=7.2043e-5: normal differences doubtful; t result is exploratory, not primary.')
    fill('M09.S4.ci95','Paired bootstrap goal-error difference [-1.471245,-0.819178] m; acceptance difference +0.74, bootstrap 95% CI [0.62,0.86], 20,000 resamples, analysis seed 2026.')
    plain(21,0,0,'Technical data verified from MuJoCo runs of the project controller; AI assistance disclosed')
    plain(21,1,0,'No benchmark outcome is generated as a sampled target; archived synthetic scaffold data excluded')
    plain(21,2,0,'Plots regenerate from actual telemetry.csv, states.npz and results.csv records')
    for key in ['prov_own_runs','prov_no_simulated','prov_plots_reproducible','prov_economics_sourced']:fill('M09.S4.'+key,'Yes')
    fill('M09.S5.linkedin_url','PENDING_SUBMISSION; no public post or student credentials have been fabricated.')
    fill('M09.S5.video_length','Original clip 01:17.4; extended obstacle/two-climb clip 02:01.4.')
    fill('M09.S5.live_mujoco','No');fill('M09.S5.business_graph','No');fill('M09.S5.all_members_speak','No')
    after(tables[22],'Supplied clips are authentic native renderings of recorded physics, not a live expert-panel capture. They contain no student webcam narration and no embedded business-sensitivity segment. Students must record those elements, post the genuine video and enter the real URL. The 22-page run guide supplies live/logged/replay instructions.')
    cases=[('A prior assistant delivery used a car-like platform although the user required the supplied M20 robot dog.','Checked the original source','M20.xml, 17 STL meshes and anatomy tests establish 12 leg plus four hub motors; the box vehicle is superseded.','Other'),
        ('The assistant ray wrapper treated mj_multiRay cutoff=5 as exact sensor range.','Official documentation (MuJoCo, standard)','The cutoff culls geoms: a plane anchored at zero disappeared far from origin. Coarse culling was disabled and actual hit distance filtered; a distant-plane regression now passes.','Wrong code / MuJoCo API'),
        ('The assistant initially used a native geometry-distance result without checking its witness segment.','Ran / tested the code or simulation','Some box queries returned zero with separated witness points. Inconsistent minima are now withheld; actual solver contact auditing remains at 2 ms.','Wrong code / MuJoCo API')]
    for i,(said,caught,truth,error) in enumerate(cases,1):
        for k,v in {'tool':'Other','said':said,'caught_by':caught,'truth':truth,'error_type':error,'verdict':'Corrected'}.items():fill(f'M09.S5.ai{i}.'+k,v)
    after(tables[23],'These are documented assistant implementation-audit cases, not invented claims that a named student personally caught each error. Personal learning, time spent, contribution share and sign-off fields remain unchanged for each member.')

    # Build only new figure/annex fragments using the original document as style source.
    annex=Document(source);ab=annex._element.body
    for node in list(ab):
        if node.tag!=qn('w:sectPr'):ab.remove(node)
    def p(text,bold=False,size=9):
        a=annex.add_paragraph();a.paragraph_format.space_after=Pt(5);a.paragraph_format.line_spacing=1.08
        run=a.add_run(text);run.font.name='Calibri';run.font.size=Pt(size);run.font.bold=bold;run.font.color.rgb=RGBColor(0,0,0);return a
    def h(text):
        a=annex.add_paragraph(style='Heading 2');a.paragraph_format.space_before=Pt(7);a.paragraph_format.space_after=Pt(5);a.paragraph_format.keep_with_next=True
        run=a.add_run(text);run.font.name='Calibri';run.font.size=Pt(11);run.font.bold=True;run.font.color.rgb=RGBColor(0,0,0)
    def page(text):
        a=annex.add_paragraph(style='Heading 1');a.paragraph_format.page_break_before=True;a.paragraph_format.keep_with_next=True
        run=a.add_run(text);run.font.name='Calibri';run.font.size=Pt(13);run.font.bold=True;run.font.color.rgb=RGBColor(0,0,0)
    def codeblock(text):
        a=p(text,size=8)
        for run in a.runs:run.font.name='Consolas'
        a.paragraph_format.keep_together=True
    def image_fragment(path,width):
        a=annex.add_paragraph();a.alignment=WD_ALIGN_PARAGRAPH.CENTER;a.paragraph_format.space_after=Pt(3)
        shape=a.add_run().add_picture(str(path),width=Inches(width));shape._inline.docPr.set('descr',path.stem.replace('_',' '));node=a._p;node.getparent().remove(node);return node
    def picture(path,caption,width=6.8):
        node=image_fragment(path,width);ab.insert(len(ab)-1,node);a=p(caption,size=8);a.paragraph_format.keep_with_next=False
    def table(headers,rows,widths=None):
        t=annex.add_table(rows=1,cols=len(headers));t.autofit=False
        widths=widths or [7.0/len(headers)]*len(headers)
        for i,width in enumerate(widths):t.columns[i].width=Inches(width)
        for i,values in enumerate([headers,*rows]):
            cells=t.rows[0].cells if i==0 else t.add_row().cells
            for j,(c,value) in enumerate(zip(cells,values)):
                c.width=Inches(widths[j]);a=c.paragraphs[0];a.paragraph_format.space_after=Pt(3);a.paragraph_format.space_before=Pt(3);a.paragraph_format.line_spacing=1.04
                run=a.add_run(str(value));run.font.name='Calibri';run.font.size=Pt(8);run.font.bold=i==0;run.font.color.rgb=RGBColor(255,255,255) if i==0 else RGBColor(0,0,0)
                tc=c._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
                for side in ['top','bottom','left','right']:
                    edge=OxmlElement('w:'+side);edge.set(qn('w:val'),'single');edge.set(qn('w:sz'),'4');edge.set(qn('w:color'),'D9D9D9');borders.append(edge)
                tc.append(borders);shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'354D5C' if i==0 else 'F1F4F6' if i%2==0 else 'FFFFFF');tc.append(shade)
                tc.append(OxmlElement('w:vAlign'));tc[-1].set(qn('w:val'),'center')
            t.rows[i]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        return t
    # Figures replace only the supplied empty picture slots; no controls are removed.
    def replace_image_slot(i,j,k,path,caption,width):
        c=cell(i,j,k)
        for old in list(c):
            if old.tag!=qn('w:tcPr'):c.remove(old)
        c.append(image_fragment(path,width));c.append(paragraph(caption,8))
    replace_image_slot(10,0,0,HERE/'dossier/free_body_diagram.png','Schematic FBD; foot normals are measured at t=2.302 s in the final seed 42 run. Sum N=338.739 N; mg=338.699 N. Approximate quasi-static agreement is not a calibrated force-sensor result.',6.6)
    cell(10,0,0).append(paragraph('Worked example: 34.525856 kg x 9.81 m/s^2 = 338.699 N. A symmetric four-foot approximation gives 84.675 N per foot; at a 0.25 m lever arm the illustrative moment is 21.169 N m. Actual per-foot normals are not assumed equal. See telemetry columns fl/fr/hl/hr_normal_N.',8))
    replace_image_slot(15,0,0,HERE/'dossier/payback_small.png','Assumed parity sensitivity. At kappa>=1 no finite positive payback is reported.',3.3)
    replace_image_slot(15,0,1,HERE/'dossier/labour_reallocation.png','Illustrative 312 h/year and 0.15 FTE; no measured time saving.',3.3)
    replace_image_slot(20,0,0,HERE/'dossier/representative_route.png','Actual seed 42 trajectory and a time-stamped online route near 30 s. This is not a fixed trajectory reference.',3.3)
    replace_image_slot(20,0,1,HERE/'dossier/goal_error_boxplot.png','Actual original benchmark: 50 runs per policy, with failures retained.',3.3)
    for i in [9,10,15,20]:
        for node in tables[i].findall('.//w:trHeight',NS):node.set(qn('w:hRule'),'atLeast')

    page('Annex A1 System architecture and evidence map')
    p('Technical implementation reviewed on 26 October 2026. This annex explains the current M20 build and the evidence behind the filled form. The main research question retains its human-workload motivation, but no human workload outcome is measured. The original six-section form and all 212 tagged controls are retained; personal declarations remain unfilled.',size=10)
    picture(HERE/'dossier/current_architecture.png','Figure A1. Implemented MuJoCo data flow. Rays and pose are ideal; camera panels are observational. All numerical outcomes derive from stepped physics.',6.8)
    table(['Evidence set','Scope and outcome','Where to verify'],[
        ('Original study','100 simulations / 50 matched environments; baseline 10/50, shared 47/50 accepted','analytics/benchmark/design.json, results.csv, statistics.json and per-run states'),
        ('Dense avoidance','26 simulations: 12 pairs plus two blocked cases; shared 12/12, baseline 0/12','analytics/obstacle_course/verification.json and statistics.json'),
        ('Two-climb course','7 simulations: six shared checks pass both rises; one illustrative baseline fails','analytics/climbing_course/verification.json and statistics.json'),
        ('Fresh extraction','31 locked packages installed separately; 44 tests pass; seed 42 state arrays bitwise equal','analytics/course_fresh_extraction_verification.json'),
        ('Dossier diagnostics','100 sampled control-effort values and six rolling-path navigation diagnostics','analytics/dossier/metrics.json and dossier_real_trials.csv')],[1.15,3.15,2.7])
    p('Paths in the form and annex are relative to the repository root unless they begin analytics/ or docs/, which are relative to Group_09_Arnav_Vihan_Pratik. The authoritative asset is M20/M20.xml and its 17 STL meshes. Source SHA256 values distinguish physical versions; paper and dossier statistics do not relabel older results as current-source experiments.')

    page('Annex A2 Mechanical model and control design')
    picture(GROUP/'docs/figures/course_climbing_frame.png','Figure A2. Recorded M20 state at 83.422 s: both rear feet are airborne during the first 8 cm terrace transfer. It is a rendered physical-state replay.',6.6)
    table(['Quantity','Implemented value','Interpretation'],[
        ('Anatomy','Four articulated legs; 12 leg and four hub motors','Supplied hybrid canine quadruped, not the earlier box vehicle'),
        ('Mass and states','34.525856 kg; nq=23, nv=22, nu=16','Supplied inertias and hierarchy; not measured hardware identification'),
        ('Leg geometry','Two sagittal links 0.25 m each; lateral offset 0.158076 m','Analytic IK checked against MuJoCo body positions'),
        ('Hub radii','Collision 0.093 m; servo conversion 0.090 m','The small mismatch is retained and disclosed; slip remains a kinematic proxy'),
        ('Physics','2 ms implicitfast; Newton solver, 100 iterations, tolerance 1e-8','Pyramidal contacts; rigid surfaces and approximate collisions'),
        ('Attitude and coordinates','Right-handed x-forward/y-left/z-up; quaternion w,x,y,z','Exact simulator state feedback; no ROS transform runtime')],[1.3,2.5,3.2])
    h('Implemented joint and stance control')
    p('tau_cmd = clip(Kp*(q_ref-q) - Kd*qdot + qfrc_bias + tau_stance, -tau_limit, +tau_limit). Leg gains and motor limits appear in Section 2. The stance term uses f = A^T*(A*A^T + 1e-4*I)^(-1)*w_des, followed by normal-force and approximate friction-box clipping. This is an instantaneous regularized allocation, not constrained MPC and not a stability proof.')
    p('The two-link IK uses cos(q_knee)=(x^2+z_eff^2-l1^2-l2^2)/(2*l1*l2), with opposite knee branches for front and rear legs. Reach is clipped for numerical validity. Leg references, wheel braking and feedback act through joint torques; the floating-base pose is set only at initialization.')
    p('For straight cruise at v=0.30 m/s, the hub angular target magnitude is v/0.09=3.333 rad/s. The code uses the negative sign for the supplied wheel-axis convention. Steering uses v_L=v-w*0.55/2 and v_R=v+w*0.55/2; this 0.55 m quantity is an effective steering parameter, not a newly measured CAD dimension.')

    page('Annex A3 Perception planning and climb state machine')
    table(['Stage','Actual input and behavior','Important boundary'],[
        ('Ray sensing','MuJoCo intersections; self geometry excluded by group; hit lengths explicitly limited to 5 m','No sensor noise, dropout, motion distortion or uncertain self-filter'),
        ('Grid','Count, sum, sum-square, mean and peak per cell; 3 cm original / 5 cm course','Exact world pose; not SLAM or camera-LiDAR fusion'),
        ('Cost','Base c=0.4*slope_norm+0.3*roughness_norm+0.3*step_norm; add obstacle proximity penalty','Base cost is clipped to [0,1]; final penalized cost need not be <=1'),
        ('Planner','8-neighbor A*, unknown cells blocked, no diagonal corner cutting; 0.70 m inflation','A measured frontier can be used while the goal remains unobserved'),
        ('Tracking','Heading P control; turn in place for error >0.30 rad; course lookahead 0.35 m','Nominal cruise 0.30 m/s; no global path or slalom waypoints preloaded'),
        ('Step edge','Broad low terrace detected; exact return x values retained in 2 cm height bands','Course sensor fixtures resolve the lip within 15 mm; no universal edge-error bound')],[1.05,3.0,2.95])
    h('Transfer sequence and gates')
    p('DOCK_ALIGN waits for |yaw|<0.025 rad and |y|<0.25 m. APPROACH moves the front feet toward the measured lip. PACK_REAR_STANCE and SHIFT_TO_REAR establish rear support; FRONT_PAIR_STEP requires fore re-contact. TORSO_RISE and BODY_ADVANCE move the body; SHIFT_TO_FRONT establishes fore support. REAR_PAIR_STEP requires rear re-contact before SETTLE. CLEAR_REAR_EDGE moves beyond the lip before nominal rolling posture resumes. A second measured terrace repeats this sequence relative to the first support level.')
    p('The final course uses a rear landing target 0.24 m beyond the measured lip, versus 0.06 m in the original single-step fixtures. These are tuned controller parameters, not a new CAD dimension. Separate lift/contact records prevent a wheel-assisted scrape or near-goal timeout being relabeled as a successful pair lift.')
    h('Delay and arbitration')
    p('LatencyQueue releases commands by simulation time at t_sent+delay, independently of frame rate. The baseline follows the delayed deterministic reactive driver. The integrated policy locally chooses a map route and transfer action, so this is an integration-level comparison; the code does not model a skilled human or isolate the causal effect of latency alone. Tilt latches at |roll|>18 degrees or |pitch|>28 degrees. No-safe-path commands stop, but a protected stop is not successful reconnaissance.')
    p('Forward rays span 180 degrees, while the closest-range decision uses the central +/-25-degree sector. Slowdown begins below 1.5 m and the local obstacle stop is 0.20 m. This differs from the course handbook 0.40 m stop example and does not establish personnel-detection or ISO conformity. A complete lattice sweep takes 0.32 simulated seconds even though partial ray integration runs at 50 Hz.')

    page('Annex A4 Experimental design and original study results')
    p('The original study uses 50 seeded matched environments, each run with both policies: 10 flat, 15 climb, 15 barrier avoidance, eight combined yard missions and two blocked scenes. Seeds are 1000-1049. Friction, low-step height, barrier position/width and rubble vary where applicable; no obstacle velocities are randomized. Deadlines are scenario-specific. Failures remain in the distributions.')
    table(['Measure','Delayed baseline','Mapped transfer policy'],[
        ('Accepted missions','10 / 50','47 / 50'),('Harmful-contact trials','30 / 50','0 / 50'),('Rollovers','0 / 50','0 / 50'),
        ('Final goal error mean +/- SD','1.3575 +/- 1.1607 m','0.2181 +/- 0.5364 m'),
        ('Simulated duration mean +/- SD','45.395 +/- 21.300 s','30.203 +/- 16.856 s'),
        ('Sampled torque-squared proxy',f"{eff['baseline']['mean']:.1f} +/- {eff['baseline']['sd']:.1f}",f"{eff['shared']['mean']:.1f} +/- {eff['shared']['sd']:.1f}")],[2.5,2.25,2.25])
    h('Primary hypothesis and matched inference')
    p('H0: the paired accepted-mission probability is equal. H1: it differs. There are 37 shared-only accepted pairs and zero baseline-only pairs. Two-sided exact McNemar p=1.4551915228e-11. Acceptance difference is +0.74, paired-bootstrap 95% CI [0.62,0.86] with 20,000 resamples and analysis seed 2026. This applies to the selected fixtures and simple baseline; significance is not evidence of industrial reliability.')
    p('For shared-minus-baseline final goal error, mean difference=-1.139460 m and difference SD=1.191647 m. Paired-bootstrap 95% CI [-1.471245,-0.819178] m. Exploratory paired t(49)=-6.761397, p=1.5540753e-8, d_z=-0.956206. Shapiro-Wilk p=7.2043e-5 argues against treating normal differences as established; Wilcoxon sensitivity p=1.0658629e-8. Do not substitute pooled Cohen d for paired d_z or call the t-test the primary binary analysis.')
    picture(HERE/'dossier/sampled_control_effort.png','Figure A3. Newly derived descriptive commanded-torque-squared proxy from actual 50 Hz states, with all failed trials retained. Units are (N m)^2 s, not J.',5.8)
    p('Original shared failures: blocked seeds 1048/1049 and a low-step combined mission at seed 1043 that reached the goal area without verifying simultaneous rear-pair liftoff. A separate long walking prototype destabilized around 24 s; the delivered braked-foot walking claim is bounded to 12 s. No universal below-three-degree roll claim is made.')

    page('Annex A5 Dense industrial course and two terrace results')
    picture(GROUP/'docs/figures/m20_obstacle_course_poster.png','Figure A4. Actual recorded seed 42 course at 30 s, with following camera, overhead scene, time-matched plan, trajectory and contact counters.',6.7)
    p('The combined course contains 20 physical hazard/boundary geometries plus two support terraces. Three 60 cm-high staggered barricades require alternating detours; drums, crates, tanks, pipe racks and rails constrain the arena. Each terrace adds 8 cm, so the final surface is 16 cm above the initial floor. Warehouse cladding, markings and bands are decorative. The M20 cruises on wheeled feet and performs articulated paired-leg climbs; it is not walking on braked feet along the full route.')
    table(['Study','Observed result','Interpretation'],[
        ('Avoidance extension','Shared 12/12 accepted; baseline 0/12; all baseline trials had harmful contact','12 matched seeds; exact paired p=0.00048828125'),
        ('Closed gates','0/2 missions accepted; 0 harmful contacts; final motion stopped','Protection is not mission success'),
        ('Combined extension','Shared 6/6 accepted with both climbs and zero harmful contacts','Six descriptive seed checks; selected static scene family'),
        ('Illustrative combined baseline','0/1 accepted; timeout with harmful contacts','One baseline is not a fully paired inferential experiment')],[1.4,3.2,2.4])
    p('The three experiment sets use different designs and source versions. They are not pooled into a larger comparative sample. Course seeds 42 and 2026-2030 vary friction and barricade offsets; both terrace rises are fixed at 8 cm. There are no participant responses, moving pedestrians or measured field outcomes.')

    page('Annex A6 Independent ascent evidence and course diagnostics')
    table(['Seed','Accepted','Time s','Goal error m','Climbs','Harmful steps','Peak pitch deg'],[
        (x['seed'],x['accepted'],f"{x['simulated_s']:.3f}",f"{x['goal_error_m']:.3f}",x['verified_climbs'],x['harmful_contact_steps'],f"{x['max_abs_pitch_deg']:.2f}") for x in course['combined_climbing'] if x['policy']=='shared'],[.55,.8,.9,1.05,.6,1.0,1.1])
    p('Seed 42 completes in 121.394 simulated seconds, with 0.107744 m goal error, 21.270557 m driven path, 350 planner calls, peak absolute roll 1.476988 degrees and pitch 12.207026 degrees. Both ascent records have front/rear pair liftoff and re-support: first verified at 87.078 s, second at 117.420 s. These values are read from the final summary and raw trajectory.')
    picture(HERE/'climbing_course/course_climbing_telemetry.png','Figure A5. Recorded torso and foot-wheel heights and foot contacts. Dashed lines mark the two independently verified ascents; contact states are directly logged.',6.6)
    p('Strict goal acceptance requires the 0.22 m goal region to persist for 0.4 s, no harmful contacts, no rollover or tilt stop and every required climb verified. Lift tests require both feet of a pair unsupported while wheel centers exceed current support height +0.13 m; later all four centers must match the raised support level +0.093 m within 0.025 m. A standing-up height increase is not accepted as climbing.')
    p(f"The six navigation-only rolling-path RMSE values average {tr['mean']:.5f} +/- {tr['sd']:.5f} m. They use nearest-segment distance to the most recent time-matched online plan. Because replanning resets the reference start near the robot, this is not independent fixed-trajectory accuracy. The baseline has no comparable saved planned path; the original 100 runs do not retain complete path history.")

    page('Annex A7 Economic assumptions failure analysis and scope')
    picture(HERE/'dossier/economic_sensitivity.png','Figure A6. Deterministic calculation from stated economic assumptions. No business inputs are measured plant data.',6.7)
    p('kappa=operating parity/reclaimed-value parity=0.22/1=0.22. T_months=12*0.30/(1-0.22)=4.615385 months. delta_FTE=312/2080=0.15. The 312 annual task-hours are an assumption, not observed labor displacement. A net parity margin 1-kappa=0.78 follows only from the chosen normalized baseline. At kappa>=1 no finite positive payback is returned; invalid or nonfinite denominators are rejected.')
    p('A queue model is not calibrated because no arrival or service-time dataset exists. A labor-reallocation assumption chart is provided instead. FMEA ratings in Section 3 are explicit engineering planning scores; the tested mitigations do not empirically estimate the occurrence or detection ratings. No risk elimination, human safety benefit, or field reliability is asserted.')
    h('Retained failure mechanisms')
    table(['Failure or gap','Observed evidence','Disposition'],[
        ('Step misalignment / rear contact','Development attempts reached the upper area but missed liftoff; one made upper-leg contact','Failed raw logs retained under development_evidence/course_tuning; exact final source snapshots retained'),
        ('Long gait instability','Longer diagonal-pair prototype failed around 24 s','Walking claim remains 12 s; no sustained-gait promise'),
        ('Source range culling','Distant plane missed when geom cutoff was treated as ray range','Returned-hit length filtering and far-origin regression fixture'),
        ('Distance witness inconsistency','Some box queries returned distance inconsistent with segment endpoints','Invalid sampled clearance minima withheld; 2 ms solver contact audit is primary'),
        ('Safety handbook gap','0.20 m stop and central decision sector differ from supplied handbook examples','No ISO compliance or personnel-detection claim; hardware review required')],[1.55,2.7,2.75])
    p('Exact simulated pose, rigid surfaces, primitive collisions and supplied inertias limit transfer to hardware. Loose debris, dust, sensing uncertainty, power/electrical dynamics, human interaction and adverse industrial atmospheres are absent. More cluttered static scenes improve the demonstration but do not remove these validity limits.')

    page('Annex A8 Reproduction verification and source lineage')
    p('Use 64-bit CPython 3.12 and requirements-lock.txt. The core release records 31 locked packages. The clean extraction installed those packages independently, passed all 44 tests and reproduced qpos, qvel, ctrl, targets and time arrays bit for bit for combined-course seed 42. The Windows launcher did not register the available runtime, so the tested setup helper accepts an explicit -Python path.')
    codeblock('powershell -ExecutionPolicy Bypass -File .\\run_demo.ps1 -Setup -Mode course_climb\npowershell -ExecutionPolicy Bypass -File .\\run_demo.ps1 -Mode course_climb -Live\npowershell -ExecutionPolicy Bypass -File .\\run_demo.ps1 -Mode course_climb -Record')
    p('If the launcher cannot find Python 3.12, add -Python "C:\\path\\to\\python.exe" to -Setup. Live mode executes native interactive physics; a normal run creates authoritative logs; Record renders saved physical states. The current combined deadline is 180 simulated seconds. Compute and rendering may take longer than real time.')
    codeblock('$Py = ".\\.venv\\Scripts\\python.exe"\n$Analytics = "Group_09_Arnav_Vihan_Pratik\\analytics"\n& $Py -m pytest -q\n& $Py "$Analytics\\verify_evidence.py"\n& $Py "$Analytics\\analyze_course.py" --folder "$Analytics\\obstacle_course"\n& $Py "$Analytics\\analyze_course.py" --folder "$Analytics\\climbing_course"\n& $Py "$Analytics\\build_dossier_evidence.py"')
    table(['Source version','Dataset and digest'],[
        ('Original study','benchmark/source_snapshot/M20 | 24a1c256abcc9a5bc3d0b38ecae3354d8a9f313ff8daf0e2b9c18daa8f7b2f96'),
        ('Dense avoidance','obstacle_course/source_snapshot/M20 | f6667155f6c018f1b0f2b70341b93c9e2ddf067a0976a078f42d32168b3a8210'),
        ('Final combined course','climbing_course/source_snapshot/M20 | e492a29fd3f4c1029773d51940e3aaeedfe68e761b94a6bb9f7542e0cf26f1b9')],[1.25,5.75])
    p('To rerun a historical physical version, use a separate extracted project and copy its snapshot Python files and M20.xml into that copy\'s M20 folder while preserving meshes. Do not overwrite old trial hashes. New course batch destinations are refused if they already exist; use a timestamped local_runs folder.')
    table(['Verification family','Evidence'],[
        ('44 unit/integration checks','Anatomy, IK/FK, rays and missing hits, grid aggregation/cost, inflation, unknown blocking, diagonal corners, latency FIFO, torque limits, tilt latch, acceptance and economics'),
        ('133 final-study logs','100 original + 26 dense avoidance + 7 combined, checked as separate designs and source versions'),
        ('Media','Three current clips decode all 1,214 frames each at 10 fps; source-state and plan hashes retained'),
        ('Documents and archive','Four-page IEEE paper; 22-page run guide; filled dossier and PDF; per-file manifest and ZIP CRC/SHA checks')],[1.7,5.3])

    page('Annex A9 Rubric crosswalk and genuine team handoff')
    p('The supplied dossier states a 20-mark capstone split (10 robotics, 5 economics, 5 showcase/viva); the repository rubric states 100 marks across four milestones. This is an evidence crosswalk, not a promise of marks or a replacement for faculty policy. The source administrative instructions and submission dates are retained.')
    table(['Repository rubric item','Marks','Evidence readiness'],[
        ('M1 onboarding and problem framing','10','Problem and variables documented; IDs corrected to approved roster; GitHub handles, actual student branch/PR histories and personal review pending'),
        ('M2 physical architecture / control / telemetry','20','Full supplied M20 model, bounded joint actuation, actual simulations and 50 Hz logged states'),
        ('M3 experiment execution / statistics / economics','20','100 original runs plus separately reported extensions; exact paired inference; assumed dimensionless economics; no psychometric measurement'),
        ('M4 paper','15','Compiled four-page IEEEtran paper, three figures, two tables, six genuine DOI references; team review pending'),
        ('M4 expert live demonstration','15','Native viewer and instructions prepared; recorded physics video supplied; actual panel execution still a student action'),
        ('M4 individual viva','15','Member-specific questions and revision cards supplied; individual understanding must be demonstrated in person'),
        ('M4 genuine Git trajectory','5','No fabricated commits. Each member must personally review and make legitimate contributions and reviews')],[2.2,.6,4.2])
    h('Human actions that documentation cannot attest')
    p('The individual learning record, hours, authored-file claims, commit share, personal AI share, signatures and optional research consent remain unchanged and unselected. Peer-review scores for Groups 10 and 03 remain unfilled because those videos were not assessed. No academic signature, participant consent, faculty approval or real post URL is invented.')
    p('The source form requires webcam contributions from every member and a business sensitivity segment. Supplied clips do not contain those personal appearances or that segment. Students must capture and narrate the material, post the genuine video and enter the actual URL; the current recorded clips remain labeled as physical-state playback.')
    h('Technical handoff files')
    p('Read README.md, docs/OBSTACLE_COURSE_REPORT.md, docs/FINAL_DELIVERY_REPORT.md and the 22-page run guide. Review M20/scene.py, m20_rosbot_mapping.py, recon_controller.py, hybrid_control.py and simulation.py. Raw data, failure records, source snapshots, optional document builder and the manifest accompany the submission. No public publishing or student-author commits were performed by this documentation task.')

    page('Annex A10 References and verified source index')
    refs=json.loads((GROUP/'docs/bibliography/verification.json').read_text())
    records=refs if isinstance(refs,list) else refs.get('references',refs.get('records',[]))
    wanted=['10.1109/LRA.2018.2849506','10.1109/IROS.2018.8594448','10.1126/scirobotics.abk2822','10.1126/scirobotics.adi9641','10.1109/TSMCC.2007.905819','10.1109/IROS.2012.6386109']
    for i,doi in enumerate(wanted,1):
        record=next(x for x in records if x['doi'].lower()==doi.lower())
        venue=record.get('venue',[]);venue=', '.join(venue) if isinstance(venue,list) else venue
        year=record.get('published',{}).get('date-parts',[[None]])[0][0]
        p(f"[{i}] {record['authors']}. {record['title']}. {venue}, {year}. DOI: {doi}")
    p('These are foundational references and metadata-verified bibliography entries, not evidence that this project reproduces their learned locomotion, MPC, uncertain localization or human studies. Local citation verification records retain the original Crossref metadata and check times.')
    h('Official technical and standards references')
    p('[7] MuJoCo API reference, ray casting and mj_geomDistance: https://mujoco.readthedocs.io/en/stable/APIreference/APIfunctions.html. Used for function semantics; project-specific failure behavior was checked by execution.')
    p('[8] ISO 3691-4:2023, public scope and operating-zone context: https://www.iso.org/standard/83545.html. Public preview reviewed; no full clause-by-clause conformity assessment or certificate is claimed.')
    p('[9] REP 103, Standard Units of Measure and Coordinate Conventions: https://reps.openrobotics.org/rep-0103/. [10] REP 105, Coordinate Frames for Mobile Platforms: https://reps.openrobotics.org/rep-0105/. Used for convention vocabulary; no runtime ROS frame tree is implemented.')
    h('Local technical evidence index')
    p('Primary: benchmark/design.json, statistics.json, every trial summary/telemetry/states. Extensions: obstacle_course/ and climbing_course/ designs, raw trials and verification receipts. Dossier: dossier/metrics.json, dossier_real_trials.csv and actual-data figures. Reproduction: course_fresh_extraction_verification.json, course_test_results.xml and course_fresh_test_results.xml. Source recovery: docs/provenance/recovered_manifest.json and archive_recovery_receipt.json. Human responsibilities: TEAM_ROSTER.json and RUBRIC_COMPLIANCE_MATRIX.md.')
    p('Original input dossier SHA256: '+INPUT_HASH+'. The input is retained for provenance and control-schema preservation. Its unsupported carried-over claims are corrected or explicitly rejected in this completed document. All 212 original controls and their option lists are retained; personal and consent subtrees remain unchanged.')

    # Merge only intentional XML changes and new media into the original package.
    temp=output.parent/'dossier_annex_working.docx';annex.save(temp)
    with zipfile.ZipFile(temp) as z:annex_raw={n:z.read(n) for n in z.namelist()}
    ar=E.fromstring(annex_raw['word/document.xml']);annex_nodes=[deepcopy(n) for n in ar.find('w:body',NS) if n.tag!=qn('w:sectPr')]
    sect=body.find('w:sectPr',NS)
    for node in annex_nodes:body.insert(list(body).index(sect),node)
    for i,node in enumerate(root.xpath('//*[local-name()="docPr"]'),1):node.set('id',str(i))
    final_controls={tag(s):s for s in root.findall('.//w:sdt',NS)}
    assert list(final_controls)==list(controls)
    assert all(canon(final_controls[t].find('w:sdtPr',NS))==v for t,v in props_before.items())
    assert all(canon(final_controls[t])==v for t,v in protected_before.items())
    # Repeat source header rows where tables span a page, without changing cells.
    for t in tables:
        if len(t.findall('w:tr',NS))>1:
            row=t.find('w:tr',NS);trpr=row.find('w:trPr',NS)
            if trpr is None:trpr=E.Element(qn('w:trPr'));row.insert(0,trpr)
            if trpr.find('w:tblHeader',NS) is None:E.SubElement(trpr,qn('w:tblHeader'))
    rel=E.fromstring(raw['word/_rels/document.xml.rels']);ids={x.get('Id') for x in rel}
    for item in E.fromstring(annex_raw['word/_rels/document.xml.rels']):
        if item.get('Id') not in ids:rel.append(item)
    ct=E.fromstring(raw['[Content_Types].xml']);existing={(x.tag,x.get('Extension'),x.get('PartName')) for x in ct}
    for item in E.fromstring(annex_raw['[Content_Types].xml']):
        key=(item.tag,item.get('Extension'),item.get('PartName'))
        if key not in existing:ct.append(item)
    core=E.fromstring(raw['docProps/core.xml']);dn={'dc':'http://purl.org/dc/elements/1.1/','cp':'http://schemas.openxmlformats.org/package/2006/metadata/core-properties'}
    for key,value in [('dc:title','M20 Group 09 Technical Dossier'),('dc:creator','MDRIIA Group 09 documentation'),('cp:lastModifiedBy','AI assisted technical documentation')]:
        node=core.find(key,dn)
        if node is not None:node.text=value
    raw['word/document.xml']=E.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)
    raw['word/_rels/document.xml.rels']=E.tostring(rel,xml_declaration=True,encoding='UTF-8',standalone=True)
    raw['[Content_Types].xml']=E.tostring(ct,xml_declaration=True,encoding='UTF-8',standalone=True)
    raw['docProps/core.xml']=E.tostring(core,xml_declaration=True,encoding='UTF-8',standalone=True)
    for name,data in annex_raw.items():
        if name.startswith('word/media/'):raw[name]=data
    with zipfile.ZipFile(output,'w',zipfile.ZIP_DEFLATED) as z:
        for name,data in raw.items():z.writestr(name,data)
    temp.unlink()
    technical_unfilled=[t for t in controls if not protected(t) and t not in filled]
    assert not technical_unfilled,technical_unfilled
    receipt={'status':'BUILT_PENDING_RENDER','input_sha256':INPUT_HASH,'output_sha256':hashlib.sha256(output.read_bytes()).hexdigest(),
        'original_tables_retained':29,'original_controls_retained':212,'technical_controls_filled':len(filled),'protected_controls_unchanged':len(protected_before),
        'protected_categories':['personal learning','peer review','student authorship attribution','signatures and consent'],
        'all_control_properties_identical':True,'original_footer_styles_customXml_preserved':True,'new_images':len([n for n in raw if n.startswith('word/media/')]),
        'filled_controls':filled,'created_utc':datetime.now(timezone.utc).isoformat()}
    (GROUP/'docs/provenance/dossier_edit_receipt.json').write_text(json.dumps(receipt,indent=2),encoding='utf-8')
    print(json.dumps({k:v for k,v in receipt.items() if k!='filled_controls'},indent=2))

def main():
    p=argparse.ArgumentParser();p.add_argument('--input',type=Path,default=GROUP/'archive/dossier_original_template.docx');p.add_argument('--out',type=Path,default=GROUP/'docs/MDRIIA_Group_09_Technical_Dossier.docx');a=p.parse_args();a.out.parent.mkdir(parents=True,exist_ok=True);build(a.input,a.out)

if __name__=='__main__':
    from build_capstone_dossier import main as current_main
    current_main()
